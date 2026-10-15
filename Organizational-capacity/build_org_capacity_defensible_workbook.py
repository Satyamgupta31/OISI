from __future__ import annotations

import re
import os
import tempfile
from collections import Counter
from pathlib import Path
import statistics
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "OISI_112_Organizational_Capacity_Verification.xlsx"
OUTPUT = HERE / "OISI_Organizational_Capacity_Final_Defensible_Audit_Corrected.xlsx"

SOURCE_SHEET = "Organizational Capacity"
SOURCE_DATA_START = 4
DATA_FIELDS = [
    "001_service_name", "008_provider_name", "010_website", "013_location",
    "014_solution_category", "019_open_code_repository_url", "017_funding_needs",
    "051_business_form", "951_business_form_other", "052_nonprofit_status",
    "053_shareholders", "054_staffing_full_time", "055_staffing_volunteer",
    "056_launch_date", "057_organizational_history", "058_affiliations",
    "059_founders", "061_primary_funding_source", "961_primary_funding_source_other",
    "076_saas_options", "077_registered_service_providers", "079_technology_readiness",
    "080_maintenance_status", "099_scoss", "verification_status", "verification_notes",
]

METRICS = [
    ("M1", "Technology readiness", "079_technology_readiness", 0.20),
    ("M2", "Maintenance status", "080_maintenance_status", 0.20),
    ("H1", "Human-resource capacity", "054_staffing_full_time + 055_staffing_volunteer", 0.20),
    ("I2", "Institutional affiliations and support", "058_affiliations", 0.20),
    ("O1", "Hosted/service-delivery capacity", "076_saas_options", 0.20),
]

COLORS = {"teal": "245B55", "dark": "243238", "light": "E8F1EF", "pale": "F4F7F6", "gold": "FCECC9", "white": "FFFFFF", "line": "D5DFDC"}


def clean(value) -> str:
    if value is None:
        return ""
    text = re.sub(r"<[^>]+>", " ", str(value))
    return re.sub(r"\s+", " ", text).strip()


def norm(value) -> str:
    return clean(value).casefold()


def stage_score(value):
    text = norm(value)
    stages = {
        "actual system proven in operational environment": 9,
        "system complete and qualified": 8,
        "technology validated in relevant environment": 5,
    }
    return stages.get(text, 0) / 9 * 100


def maintenance_score(value):
    text = norm(value)
    if text == "actively maintained":
        return 100
    if text == "minimally maintained":
        return 50
    return 0


def staffing_score(value):
    text = norm(value)
    return {"0": 0, "1-5": 25, "6-10": 50, "11-20": 75, "more than 20": 100}.get(text, 0)


def staffing_observed(value):
    return norm(value) in {"0", "1-5", "6-10", "11-20", "more than 20"}


def affiliation_score(value):
    text = norm(value)
    if not text:
        return 0, "Blank: zero under the agreed unavailable-data rule"
    placeholders = {"n/a", "na", "none", "no", "not applicable", "unknown", "see website", "no affiliation", "no affiliations", "no institutional affiliation"}
    if text in placeholders or re.fullmatch(r"(?:no|none|not applicable|n/?a|unknown)(?:\s+affiliations?)?[.! ]*", text):
        return 0, "Explicit absence or placeholder"
    support_terms = ("supported by", "support from", "hosted by", "organizational home", "organizational host", "fiscal sponsor", "fiscal sponsorship", "operational support", "financial support", "funded by", "managing day-to-day", "manages day-to-day", "administratively hosted")
    negation_terms = ("no support", "without support", "not supported", "no direct involvement", "no affiliation", "not affiliated")
    if any(term in text for term in support_terms) and not any(term in text for term in negation_terms):
        return 100, "Explicit host or operational/financial support language"
    return 50, "Substantive affiliation or membership reported; level of operational support not established"


def affiliation_observed(value):
    text = norm(value)
    return bool(text) and text not in {"n/a", "na", "unknown", "see website", "not known", "not reported"}


def service_score(value):
    text = norm(value)
    if text.startswith("through "):
        return 100
    return 0


def service_observed(value):
    text = norm(value)
    return text == "no services available" or text.startswith("through ")


def style_sheet(ws, widths=None, freeze="A2"):
    ws.freeze_panes = freeze
    if ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=COLORS["teal"])
        cell.font = Font(color=COLORS["white"], bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        if widths and col in widths:
            width = widths[col]
        else:
            samples = [len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 100) + 1)]
            width = min(max(max(samples, default=10) + 2, 12), 48)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32


def read_records():
    wb = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)
    source_ws = wb[SOURCE_SHEET]
    headers = [source_ws.cell(2, col).value for col in range(1, source_ws.max_column + 1)]
    positions = {header: i + 1 for i, header in enumerate(headers)}
    records = []
    for row_num in range(SOURCE_DATA_START, source_ws.max_row + 1):
        name = clean(source_ws.cell(row_num, positions["001_service_name"]).value)
        if not name:
            continue
        records.append({key: source_ws.cell(row_num, positions[key]).value for key in DATA_FIELDS})
    return wb, records


def write_formula_caches(path: Path, caches: dict[str, dict[str, float]]):
    """Store calculated numeric values beside formulas for previewers that do not recalculate."""
    ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
          "pkg": "http://schemas.openxmlformats.org/package/2006/relationships"}
    ET.register_namespace("", ns["main"])
    ET.register_namespace("r", ns["rel"])
    temp_name = None
    with ZipFile(path, "r") as source_zip:
        workbook_root = ET.fromstring(source_zip.read("xl/workbook.xml"))
        rel_root = ET.fromstring(source_zip.read("xl/_rels/workbook.xml.rels"))
        targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rel_root.findall("pkg:Relationship", ns)}
        sheet_paths = {}
        for sheet in workbook_root.findall("main:sheets/main:sheet", ns):
            target = targets[sheet.attrib[f"{{{ns['rel']}}}id"]]
            sheet_paths[sheet.attrib["name"]] = target.lstrip("/") if target.startswith("/") else f"xl/{target}"

        replacements = {}
        for sheet_name, values in caches.items():
            xml_path = sheet_paths[sheet_name]
            root = ET.fromstring(source_zip.read(xml_path))
            cells = {cell.attrib.get("r"): cell for cell in root.findall(".//main:c", ns)}
            for address, value in values.items():
                cell = cells[address]
                if cell.find("main:f", ns) is None:
                    raise ValueError(f"Expected formula at {sheet_name}!{address}")
                value_node = cell.find("main:v", ns)
                if value_node is None:
                    value_node = ET.SubElement(cell, f"{{{ns['main']}}}v")
                value_node.text = format(float(value), ".15g")
                cell.attrib.pop("t", None)
            replacements[xml_path] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

        fd, temp_name = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
        os.close(fd)
        with ZipFile(temp_name, "w", ZIP_DEFLATED) as target_zip:
            for item in source_zip.infolist():
                target_zip.writestr(item, replacements.get(item.filename, source_zip.read(item.filename)))
    try:
        os.replace(temp_name, path)
    finally:
        if temp_name and os.path.exists(temp_name):
            os.remove(temp_name)


def compute(row):
    m1 = stage_score(row["079_technology_readiness"])
    m2 = maintenance_score(row["080_maintenance_status"])
    ft = staffing_score(row["054_staffing_full_time"])
    volunteer = staffing_score(row["055_staffing_volunteer"])
    h1 = (ft + volunteer) / 2
    i2, i2_reason = affiliation_score(row["058_affiliations"])
    o1 = service_score(row["076_saas_options"])
    total = (m1 + m2 + h1 + i2 + o1) / 5
    observed = {
        "M1": norm(row["079_technology_readiness"]) in {
            "actual system proven in operational environment", "system complete and qualified",
            "technology validated in relevant environment",
        },
        "M2": norm(row["080_maintenance_status"]) in {
            "actively maintained", "minimally maintained", "inactive", "not maintained",
        },
        "FT": staffing_observed(row["054_staffing_full_time"]),
        "Volunteer": staffing_observed(row["055_staffing_volunteer"]),
        "I2": affiliation_observed(row["058_affiliations"]),
        "O1": service_observed(row["076_saas_options"]),
    }
    observed_parts = [score for score, is_present in ((ft, observed["FT"]), (volunteer, observed["Volunteer"])) if is_present]
    observed_h1 = sum(observed_parts) / len(observed_parts) if observed_parts else 0
    metric_observed = {"M1": observed["M1"], "M2": observed["M2"], "H1": observed["FT"] or observed["Volunteer"], "I2": observed["I2"], "O1": observed["O1"]}
    available_scores = [score for score, present in (
        (m1, metric_observed["M1"]), (m2, metric_observed["M2"]), (observed_h1, metric_observed["H1"]),
        (i2, metric_observed["I2"]), (o1, metric_observed["O1"]),
    ) if present]
    available_count = sum(metric_observed.values())
    observed_mean = sum(available_scores) / len(available_scores) if available_scores else 0
    flags = []
    if not clean(row["054_staffing_full_time"]):
        flags.append("Full-time staffing missing/unknown scored 0")
    if not clean(row["055_staffing_volunteer"]):
        flags.append("Volunteer staffing missing/unknown scored 0")
    if clean(row["058_affiliations"]):
        flags.append("Review affiliation coding; override available")
    if norm(row["076_saas_options"]) in {"not applicable", "unknown"}:
        flags.append("Service-delivery field not applicable/unknown scored 0")
    return {"M1": m1, "M2": m2, "FT": ft, "Volunteer": volunteer, "H1": h1, "I2": i2, "I2_reason": i2_reason, "O1": o1, "Total": total, "Flags": "; ".join(flags), "Observed": observed, "Metric observed": metric_observed, "Observed H1": observed_h1, "Available count": available_count, "Coverage": available_count / 5, "Observed mean": observed_mean}


def main():
    source_wb, records = read_records()
    computed = [compute(row) for row in records]
    n = len(records)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    names = ["Summary", "Scoring", "Metric Rules", "Methodology", "Parameters", "Verification", "Affiliation Coding", "Coverage", "Diagnostics", "Calc", "Data", "Field Guide"]
    sh = {name: wb.create_sheet(name) for name in names}

    rules = sh["Metric Rules"]
    rules.append(["Code", "Metric", "Source field(s)", "Weight", "Scoring criteria", "Unavailable-data treatment", "Interpretation / caution"])
    rule_rows = [
        ["M1", "Technology readiness", "079", 0.20, "Nine-stage source scale: score = level / 9 x 100. Observed: Actual system proven in operational environment=100; System complete and qualified=88.89; Technology validated in relevant environment=55.56.", "Unknown, blank, or unrecognized=0.", "Ordinal normalization on the source readiness scale; current data are concentrated at the highest level."],
        ["M2", "Maintenance status", "080", 0.20, "Actively Maintained=100; Minimally Maintained=50; Inactive/not maintained=0.", "Unknown or blank=0.", "Current responses are highly concentrated at actively maintained."],
        ["H1", "Human-resource capacity", "054 + 055", 0.20, "Score each field: 0 people=0; 1-5=25; 6-10=50; 11-20=75; >20=100. H1=(full-time score+volunteer score)/2.", "Unknown or blank in either field=0 for that field; H1 remains the average of both fields.", "Ordinal ranges, not exact headcounts. Paid and volunteer staffing each contribute half of H1."],
        ["I2", "Institutional affiliations and support", "058", 0.20, "Formal host or explicit operational/financial support=100; substantive affiliation/membership without stated operational support=50; explicit absence/placeholders=0.", "Blank or unknown=0. Free-text provisional coding can be overridden on Verification sheet with 0, 50, or 100.", "Narrative coding needs consistent review; an override changes the Scoring sheet formula after recalculation in Excel."],
        ["O1", "Hosted/service-delivery capacity", "076", 0.20, "Any 'Through ...' hosted/service route=100; No services available=0.", "Not applicable, unknown, or blank=0.", "Measures reported service availability, not quality, uptime, affordability, or control."],
    ]
    for row in rule_rows:
        rules.append(row)

    method = sh["Methodology"]
    method.append(["Methodology and interpretation", "Details"])
    methodology = [
        ["Metric selection", "Five measures are retained: technology readiness, maintenance, combined staffing, institutional affiliations/support, and hosted/service-delivery capacity. M3 Infrastructure age and I1 Formal organizational host are excluded. Funding, legal form, SCOSS, and provider-list fields remain descriptive, not scored."],
        ["Equal weighting", "Each of the five metrics receives 20%. Equal weighting is a transparent analytic choice; it is not estimated from the observed data or claimed to establish equal causal importance."],
        ["Final formula", "Organizational Capacity Score = (M1 + M2 + H1 + I2 + O1) / 5. Each metric is first placed on a 0-100 scale."],
        ["Missing-data rule", "No coverage or eligibility threshold is applied. Blank, unknown, not applicable, or unrecognized values score 0, as requested. The denominator remains five for every infrastructure."],
        ["Observed coverage diagnostics", "Observed coverage is reported for transparency only; it is not an eligibility gate. A metric is observed when a valid source response exists. H1 is observed if at least one staffing subfield has a recognized category; its observed-only diagnostic averages only the staffing subfields actually reported, while the final H1 score assigns zero to a missing subfield."],
        ["Observed-only mean", "The per-infrastructure observed-only mean averages available metric scores and excludes unavailable metrics. It uses a variable denominator and is descriptive; the final score remains the equal-weight zero-filled mean across all five metrics."],
        ["Interpretation", "The zero-filled score reflects recorded capacity and data availability together. A zero for a missing field does not prove the underlying capacity is absent. Publish metric-level values and data-quality notes with the total."],
        ["Human resources", "Full-time and volunteer staffing are subfields of one metric. Each is scored 0/25/50/75/100 by reported category; their two sub-scores are averaged. Missing/unknown subfields score zero. The category scale is ordinal."],
        ["Affiliation coding", "Field 058 is free text. A provisional deterministic rule assigns 100 only where explicit hosting/support language is present; other substantive affiliation text receives 50; explicit absence/placeholders/blank receive 0. Review and numeric overrides (0, 50, 100) are available in Verification."],
        ["Source data", f"The workbook contains {n} source entries copied from the Organizational Capacity verification sheet and includes the source Field_Guide. The source data are preserved as a separate sheet."],
        ["No performance bands", "No 70-point cutoff or high/medium/low categories are applied."],
        ["Limitations", "Technology readiness and maintenance responses show little variation. Staffing and affiliation information is incomplete or narrative. Hosted/service delivery is not a direct measure of quality. These constraints should be stated when interpreting comparisons."],
    ]
    for row in methodology:
        method.append(row)

    params = sh["Parameters"]
    params.append(["Code", "Metric", "Source field(s)", "Exact weight", "Displayed weight", "Scale"])
    for code, title, fields, weight in METRICS:
        params.append([code, title, fields, weight, "20.00%", "0-100"])
    params.append(["Total", "", "", 1.0, "100.00%", ""])
    params.append([])
    params.append(["Zero-fill policy", "All unavailable, unknown, not applicable, or unrecognized values score 0; denominator stays five."])
    params.append(["Coverage threshold", "None"])

    # Keep the source values intact and expose narrative-affiliation coding for review.
    data = sh["Data"]
    data.append(DATA_FIELDS)
    for row in records:
        data.append([row.get(field) for field in DATA_FIELDS])

    guide = sh["Field Guide"]
    src_guide = source_wb["Field_Guide"]
    for row in src_guide.iter_rows(values_only=True):
        guide.append(list(row))

    verification = sh["Verification"]
    verification.append(["Service", "Raw 058 affiliation text", "Provisional I2 score", "Override I2 (0/50/100)", "Final I2 score", "Reviewer notes"])
    affiliation = sh["Affiliation Coding"]
    affiliation.append(["Service", "Raw 058 affiliation text", "Provisional score", "Coding rationale", "Override score", "Final score", "Review note"])
    calc = sh["Calc"]
    calc.append(["Service", "M1 Technology readiness", "M2 Maintenance", "Full-time staffing subscore", "Volunteer staffing subscore", "H1 Human resources", "I2 Provisional affiliations", "I2 rationale", "O1 Service delivery", "Python-checked total", "Scoring flags"])
    scoring = sh["Scoring"]
    scoring.append(["Service", "Provider", "Solution category", "079 Readiness", "080 Maintenance", "054 Full-time staff", "055 Volunteers", "058 Affiliations", "076 Service options", "M1", "M2", "H1", "I2", "O1", "Final score / 100", "Rank", "Data-quality notes", "Available metrics (of 5)", "Observed coverage", "Observed-only mean"])
    coverage_ws = sh["Coverage"]
    coverage_ws.append(["Service", "M1 observed", "M2 observed", "054 staff observed", "055 volunteer observed", "H1 observed (either field)", "I2 observed", "O1 observed", "Available metrics (of 5)", "Observed coverage", "Observed-only mean", "Zero-filled final score"])
    diagnostic_ws = sh["Diagnostics"]
    diagnostic_ws.append(["Code", "Metric", "Source field(s)", "Observed n", "Observed coverage", "Mean among observed", "Mean under zero-fill", "SD among observed", "% zero among observed", "Coverage definition / reading"])

    for idx, (row, values) in enumerate(zip(records, computed), start=2):
        service = clean(row["001_service_name"])
        raw_aff = row["058_affiliations"]
        verification.append([service, raw_aff, values["I2"], None, f'=IF(ISNUMBER(D{idx}),D{idx},C{idx})', "Enter 0, 50, or 100 only when correcting the provisional text coding."])
        affiliation.append([service, raw_aff, values["I2"], values["I2_reason"], None, f"=Verification!E{idx}", "Use the override in Verification; this sheet mirrors the active score."])
        calc.append([service, values["M1"], values["M2"], values["FT"], values["Volunteer"], values["H1"], values["I2"], values["I2_reason"], values["O1"], values["Total"], values["Flags"]])
        source_row = idx + SOURCE_DATA_START - 2
        scoring.append([
            service, clean(row["008_provider_name"]), clean(row["014_solution_category"]),
            clean(row["079_technology_readiness"]), clean(row["080_maintenance_status"]),
            clean(row["054_staffing_full_time"]), clean(row["055_staffing_volunteer"]),
            clean(raw_aff), clean(row["076_saas_options"]),
            f"=Calc!B{idx}", f"=Calc!C{idx}", f"=Calc!F{idx}", f"=Verification!E{idx}", f"=Calc!I{idx}",
            f"=AVERAGE(J{idx}:N{idx})",
            f'=RANK(O{idx},$O$2:$O${n+1},0)+(COUNTIF($O$2:$O${n+1},O{idx})-1)/2',
            values["Flags"],
            f"=Coverage!I{idx}", f"=Coverage!J{idx}", f"=Coverage!K{idx}"
        ])
        obs = values["Observed"]
        metric_obs = values["Metric observed"]
        coverage_ws.append([
            service, int(obs["M1"]), int(obs["M2"]), int(obs["FT"]), int(obs["Volunteer"]),
            int(metric_obs["H1"]), int(obs["I2"]), int(obs["O1"]), values["Available count"],
            values["Coverage"], values["Observed mean"], values["Total"],
        ])

    metric_definitions = [
        ("M1", "Technology readiness", "079", "A recognized readiness category is present."),
        ("M2", "Maintenance status", "080", "A recognized active/minimal/inactive/not-maintained response is present."),
        ("H1", "Human-resource capacity", "054 + 055", "At least one staffing field has a recognized range; observed-only H1 averages only reported staffing fields."),
        ("I2", "Institutional affiliations and support", "058", "A substantive affiliation or explicit no-affiliation response is present; blank, unknown, and placeholders are unavailable."),
        ("O1", "Hosted/service-delivery capacity", "076", "A hosted route or explicit 'No services available' response is present; N/A and Unknown are unavailable."),
    ]
    for code, title, fields, definition in metric_definitions:
        if code == "H1":
            observed_values = [item["Observed H1"] for item in computed if item["Metric observed"]["H1"]]
            zero_values = [item["H1"] for item in computed]
        else:
            observed_values = [item[code] for item in computed if item["Metric observed"][code]]
            zero_values = [item[code] for item in computed]
        count = len(observed_values)
        diagnostic_ws.append([
            code, title, fields, count, count / n,
            sum(observed_values) / count if count else 0,
            sum(zero_values) / n,
            statistics.pstdev(observed_values) if count > 1 else 0,
            sum(value == 0 for value in observed_values) / count if count else 0,
            definition,
        ])
    diagnostic_ws.append([])
    diagnostic_ws.append(["Interpretation", "Observed coverage is descriptive only and does not determine whether an entry receives a score."])
    diagnostic_ws.append(["Zero-filled mean", "Mean under zero-fill uses the final scoring rule, including zero for unavailable values and a fixed denominator."])
    diagnostic_ws.append(["Observed-only mean", "Excludes unavailable metrics and has a variable denominator; it is not the final equal-weight score."])

    summary = sh["Summary"]
    summary.append(["Organizational Capacity: Final Equal-Weight Results", ""])
    summary.append(["Infrastructure entries", n])
    summary.append(["Final metrics", "M1, M2, H1, I2, O1"])
    summary.append(["Weight per metric", "20.00% (1/5)"])
    summary.append(["Coverage / eligibility gate", "None"])
    summary.append(["Missing, unknown, N/A", "Scored 0; denominator remains five"])
    summary.append(["Mean score", f"=AVERAGE(Scoring!O2:O{n+1})"])
    summary.append(["Median score", f"=MEDIAN(Scoring!O2:O{n+1})"])
    summary.append(["Highest score", f"=MAX(Scoring!O2:O{n+1})"])
    summary.append(["Lowest score", f"=MIN(Scoring!O2:O{n+1})"])
    summary.append(["Important interpretation", "Zero-filled score reflects both reported capacity and missing information. Review metric-level scores and notes before interpreting rank differences."])
    summary.append(["Affiliation review", "Field 058 is free text. Use Verification to enter a corrected I2 score of 0, 50, or 100; scoring formulas update when opened/recalculated in Excel."])
    summary.append(["Mean observed coverage", f"=AVERAGE(Coverage!J2:J{n+1})"])
    summary.append(["Entries with all five metrics observed", f"=COUNTIF(Coverage!I2:I{n+1},5)"])

    # Formatting and number formats.
    for name, ws in sh.items():
        style_sheet(ws)
    style_sheet(summary, {1: 32, 2: 105}, freeze="A2")
    style_sheet(rules, {1: 10, 2: 32, 3: 30, 4: 12, 5: 72, 6: 58, 7: 62})
    style_sheet(method, {1: 27, 2: 115})
    style_sheet(params, {1: 14, 2: 36, 3: 52, 4: 15, 5: 18, 6: 18})
    style_sheet(scoring, {1: 35, 2: 28, 3: 32, 4: 27, 5: 23, 6: 19, 7: 17, 8: 55, 9: 50, 10: 14, 11: 14, 12: 14, 13: 14, 14: 14, 15: 18, 16: 10, 17: 62, 18: 21, 19: 18, 20: 20})
    style_sheet(verification, {1: 35, 2: 70, 3: 18, 4: 22, 5: 18, 6: 68})
    style_sheet(affiliation, {1: 35, 2: 70, 3: 18, 4: 66, 5: 18, 6: 18, 7: 58})
    style_sheet(calc, {1: 35, 2: 20, 3: 18, 4: 24, 5: 23, 6: 18, 7: 22, 8: 66, 9: 20, 10: 22, 11: 68})
    style_sheet(coverage_ws, {1: 35, 2: 15, 3: 15, 4: 18, 5: 20, 6: 24, 7: 15, 8: 15, 9: 22, 10: 18, 11: 20, 12: 22})
    style_sheet(diagnostic_ws, {1: 12, 2: 36, 3: 30, 4: 15, 5: 18, 6: 22, 7: 22, 8: 20, 9: 23, 10: 82})
    style_sheet(data, {i + 1: min(max(len(field) + 2, 16), 38) for i, field in enumerate(DATA_FIELDS)})
    style_sheet(guide, {1: 38, 2: 32, 3: 64, 4: 72})

    for ws_name in ("Scoring", "Calc"):
        ws = sh[ws_name]
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (float, int)) and cell.column not in (1,):
                    cell.number_format = "0.00"
    for row in range(2, n + 2):
        for col in range(10, 16):
            scoring.cell(row, col).number_format = "0.00"
        for col in (3, 5, 6):
            verification.cell(row, col).number_format = "0.00"
        for col in (3, 5, 6):
            affiliation.cell(row, col).number_format = "0.00"
        coverage_ws.cell(row, 10).number_format = "0.00%"
        coverage_ws.cell(row, 11).number_format = "0.00"
        coverage_ws.cell(row, 12).number_format = "0.00"
        scoring.cell(row, 15).number_format = "0.00"
        scoring.cell(row, 19).number_format = "0.00%"
        scoring.cell(row, 20).number_format = "0.00"
    for row in range(2, 7):
        diagnostic_ws.cell(row, 5).number_format = "0.00%"
        diagnostic_ws.cell(row, 6).number_format = "0.00"
        diagnostic_ws.cell(row, 7).number_format = "0.00"
        diagnostic_ws.cell(row, 8).number_format = "0.00"
        diagnostic_ws.cell(row, 9).number_format = "0.00%"

    ws = sh["Verification"]
    for row in range(2, n + 2):
        cell = ws.cell(row, 4)
        cell.fill = PatternFill("solid", fgColor=COLORS["gold"])
        cell.comment = openpyxl.comments.Comment("Optional reviewer override. Allowed values: 0, 50, or 100.", "OpenAI")

    sh["Summary"].freeze_panes = "A2"
    sh["Summary"].auto_filter.ref = "A1:B14"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)
    totals = [item["Total"] for item in computed]
    ordered = sorted(totals, reverse=True)
    formula_caches = {
        "Summary": {
            "B7": sum(totals) / n,
            "B8": (ordered[(n - 1) // 2] + ordered[n // 2]) / 2,
            "B9": max(totals),
            "B10": min(totals),
            "B13": sum(item["Coverage"] for item in computed) / n,
            "B14": sum(item["Available count"] == 5 for item in computed),
        },
        "Scoring": {},
        "Verification": {},
        "Affiliation Coding": {},
    }
    for excel_row, values in enumerate(computed, start=2):
        for col, key in zip("JKLMN", ("M1", "M2", "H1", "I2", "O1")):
            formula_caches["Scoring"][f"{col}{excel_row}"] = values[key]
        formula_caches["Scoring"][f"O{excel_row}"] = values["Total"]
        greater = sum(other > values["Total"] for other in totals)
        tied = sum(other == values["Total"] for other in totals)
        formula_caches["Scoring"][f"P{excel_row}"] = greater + 1 + (tied - 1) / 2
        formula_caches["Scoring"][f"R{excel_row}"] = values["Available count"]
        formula_caches["Scoring"][f"S{excel_row}"] = values["Coverage"]
        formula_caches["Scoring"][f"T{excel_row}"] = values["Observed mean"]
        formula_caches["Verification"][f"E{excel_row}"] = values["I2"]
        formula_caches["Affiliation Coding"][f"F{excel_row}"] = values["I2"]
    write_formula_caches(OUTPUT, formula_caches)
    print(f"Wrote {OUTPUT} with {n} entries and {len(METRICS)} equal-weight metrics.")
    print(f"Python-checked mean: {sum(x['Total'] for x in computed)/n:.4f}")
    print(f"Affiliation nonblank: {sum(bool(clean(r['058_affiliations'])) for r in records)}")


if __name__ == "__main__":
    main()
