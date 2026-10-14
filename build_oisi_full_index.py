from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from xml.sax.saxutils import escape

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT_XLSX = Path("OISI_Full_Index_All_112.xlsx")
OUT_DOCX = Path("OISI_Full_Index_All_112_Summary.docx")


def norm(value) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def num(value):
    if value in (None, "", "NA", "None"):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_table(path, sheet, header_row=1):
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = {h: ws.cell(r, i + 1).value for i, h in enumerate(headers) if h}
        if any(v is not None for v in row.values()):
            rows.append(row)
    return rows


def by_name(rows, name_key):
    return {norm(row.get(name_key)): row for row in rows}


def status_from_score(score, status=None):
    if status:
        return status
    return "Scored" if score is not None else "Not scored"


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(10, min(55, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[letter].width = width


def add_table(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    style_sheet(ws)


def build_rows():
    ce_data = read_table("Community-Engagement/OISI_Community_Engagement_v3_multi_sheet.xlsx", "Data")
    ce_scores = by_name(read_table("Community-Engagement/OISI_Community_Engagement_v3_multi_sheet.xlsx", "Scoring"), "Service")
    gov_scores = by_name(read_table("Governance-Practices/OISI_Governance_Practices_Scoring.xlsx", "Scores"), "Service")
    org_scores = by_name(read_table("Organizational-capacity/OISI_Organizational_Capacity_Scoring_v3.xlsx", "Scores"), "Service")
    dev_scores = by_name(read_table("Development-Activity/Dev Data/oisi_c4/OISI_Component4_Final_Master_With_Provenance.xlsx", "Scoring"), "infrastructure")

    rows = []
    for idx, base in enumerate(ce_data, 1):
        key = norm(base["001_service_name"])
        ce = ce_scores.get(key, {})
        gov = gov_scores.get(key, {})
        org = org_scores.get(key, {})
        dev = dev_scores.get(key, {})

        ce_score = num(ce.get("COMMUNITY ENGAGEMENT SCORE (0-100)"))
        gov_score = num(gov.get("Governance score (0-100)"))
        org_score = num(org.get("ORG CAPACITY SCORE (0-100)"))
        dev_score = num(dev.get("component4_score"))
        component_scores = [ce_score, gov_score, org_score, dev_score]
        available = [score for score in component_scores if score is not None]
        available_count = len(available)
        strict_score = sum(component_scores) / 4 if available_count == 4 else None
        available_score = sum(available) / available_count if available else None
        evidence_adjusted_score = sum(score or 0 for score in component_scores) / 4

        missing = []
        if ce_score is None:
            missing.append("Community Engagement")
        if gov_score is None:
            missing.append("Governance Practices")
        if org_score is None:
            missing.append("Organizational Capacity")
        if dev_score is None:
            missing.append("Development Activity")

        missing_notes = []
        missing_statuses = [
            ("Community Engagement", ce.get("Result status")),
            ("Governance Practices", "No component score; see completeness/evidence confidence"),
            ("Organizational Capacity", org.get("Result status")),
            ("Development Activity", dev.get("score_status") or dev.get("master_record_status") or dev.get("classification")),
        ]
        for component, reason in missing_statuses:
            if component in missing:
                missing_notes.append(f"{component}: {reason or 'no source status recorded'}")

        if available_count == 4:
            index_status = "Complete OISI index"
        elif available_count >= 2:
            index_status = "Partial OISI index"
        else:
            index_status = "Insufficient component data"

        rows.append(
            {
                "Index row": idx,
                "Service": base.get("001_service_name"),
                "Provider": base.get("008_provider_name"),
                "Website": base.get("010_website"),
                "Location": base.get("013_location"),
                "Solution category": base.get("014_solution_category"),
                "Repository URL": base.get("019_open_code_repository_url"),
                "OISI score strict all 4 (0-100)": round(strict_score, 2) if strict_score is not None else "",
                "Evidence-adjusted OISI (missing=0) (0-100)": round(evidence_adjusted_score, 2),
                "OISI score over available components (0-100)": round(available_score, 2) if available_score is not None else "",
                "Components available (of 4)": available_count,
                "Index status": index_status,
                "Missing components": "; ".join(missing),
                "Missing score notes": "; ".join(missing_notes),
                "Community Engagement score": round(ce_score, 2) if ce_score is not None else "",
                "Community Engagement status": ce.get("Result status") or "",
                "Governance Practices score": round(gov_score, 2) if gov_score is not None else "",
                "Governance completeness": gov.get("Completeness") if gov.get("Completeness") is not None else "",
                "Governance evidence confidence": gov.get("Evidence confidence") if gov.get("Evidence confidence") is not None else "",
                "Organizational Capacity score": round(org_score, 2) if org_score is not None else "",
                "Organizational Capacity status": org.get("Result status") or "",
                "Development Activity score": round(dev_score, 2) if dev_score is not None else "",
                "Development Activity status": dev.get("score_status") or dev.get("master_record_status") or "",
                "Development classification": dev.get("classification") or "",
                "Development repo used": dev.get("repo") or "",
                "Development host": dev.get("host") or "",
                "Development available weight %": dev.get("available_weight_pct") or "",
                "Development issue tracker": dev.get("issue_tracker") or "",
                "Content summary": (
                    f"{base.get('001_service_name')} is a {base.get('014_solution_category')} provided by "
                    f"{base.get('008_provider_name')}. Repository/source used for development activity: "
                    f"{dev.get('repo') or base.get('019_open_code_repository_url') or 'not scoreable / not available'}."
                ),
            }
        )
    ranked = sorted([r for r in rows if r["OISI score strict all 4 (0-100)"] != ""], key=lambda r: r["OISI score strict all 4 (0-100)"], reverse=True)
    rank = {r["Service"]: i for i, r in enumerate(ranked, 1)}
    for r in rows:
        r["Strict OISI rank"] = rank.get(r["Service"], "")
    adjusted_ranked = sorted(rows, key=lambda r: r["Evidence-adjusted OISI (missing=0) (0-100)"], reverse=True)
    adjusted_rank = {r["Service"]: i for i, r in enumerate(adjusted_ranked, 1)}
    for r in rows:
        r["Evidence-adjusted OISI rank"] = adjusted_rank[r["Service"]]
    return rows


def governance_score_blank_sensitivity(record, blank_as_zero: bool):
    """Recalculate Governance with blank 050/048/046 values as zero or missing."""
    status_values = {
        "implemented/available": 1.0,
        "in progress": 0.67,
        "considering": 0.33,
        "not planning": 0.0,
        "no direct costs": 1.0,
    }

    def status(field):
        value = record.get(field)
        return status_values.get(str(value).strip().lower()) if value not in (None, "") else None

    blank_value = 0.0 if blank_as_zero else None
    board = record.get("050 Board structure")
    s048_raw = record.get("048 Community governance")
    s048 = {"formal": 1.0, "ad hoc": 0.5, "none": 0.0}.get(str(s048_raw).strip().lower()) if s048_raw not in (None, "") else blank_value
    frameworks = record.get("046 Values frameworks")
    s046 = 1.0 if frameworks not in (None, "") else blank_value

    formal_items = [
        (12, status("024 Governance structure")),
        (12, status("026 Governance records")),
        (8, status("108 Bylaws")),
        (8, 1.0 if board not in (None, "") else blank_value),
    ]
    participatory_items = [(20, s048)]
    transparency_items = [
        (7, status("029 Web accessibility")),
        (7, status("033 Transparent pricing")),
        (7, status("035 Equity commitment")),
        (7, status("044 Privacy commitment")),
        (7, s046),
    ]

    def weighted_mean(items):
        observed = [(weight, value) for weight, value in items if value is not None]
        if not observed:
            return None
        return sum(weight * value for weight, value in observed) / sum(weight for weight, _ in observed)

    formal = weighted_mean(formal_items)
    participatory = weighted_mean(participatory_items)
    transparency = weighted_mean(transparency_items)
    observed_weight = sum(weight for items in [formal_items, participatory_items, transparency_items] for weight, value in items if value is not None)
    completeness = observed_weight / 95
    if completeness < 0.70 or formal is None or transparency is None:
        return None
    subcomponents = [(0.40, formal), (0.25, participatory), (0.35, transparency)]
    observed = [(weight, value) for weight, value in subcomponents if value is not None]
    return 100 * sum(weight * value for weight, value in observed) / sum(weight for weight, _ in observed)


def rank_map(values):
    ordered = sorted(values.items(), key=lambda item: item[1])
    ranks = {}
    start = 0
    while start < len(ordered):
        end = start + 1
        while end < len(ordered) and ordered[end][1] == ordered[start][1]:
            end += 1
        average_rank = ((start + 1) + end) / 2
        for name, _ in ordered[start:end]:
            ranks[name] = average_rank
        start = end
    return ranks


def rank_comparison(reference, candidate):
    common = sorted(set(reference) & set(candidate))
    if len(common) < 2:
        return len(common), "", ""
    a = rank_map({name: reference[name] for name in common})
    b = rank_map({name: candidate[name] for name in common})
    av = [a[name] for name in common]
    bv = [b[name] for name in common]
    ma, mb = sum(av) / len(av), sum(bv) / len(bv)
    denom = (sum((x - ma) ** 2 for x in av) * sum((y - mb) ** 2 for y in bv)) ** 0.5
    rho = sum((x - ma) * (y - mb) for x, y in zip(av, bv)) / denom if denom else 1.0
    common_reference = {name: reference[name] for name in common}
    common_candidate = {name: candidate[name] for name in common}
    top_a = set(sorted(common_reference, key=common_reference.get, reverse=True)[:10])
    top_b = set(sorted(common_candidate, key=common_candidate.get, reverse=True)[:10])
    return len(common), round(rho, 4), len(top_a & top_b)


def build_sensitivity(rows):
    gov_rows = by_name(read_table("Governance-Practices/OISI_Governance_Practices_Scoring.xlsx", "Scores"), "Service")
    complete = [r for r in rows if r["Components available (of 4)"] == 4]
    component_names = ["Community Engagement", "Governance Practices", "Organizational Capacity", "Development Activity"]
    score_columns = [f"{name} score" for name in component_names]
    complete_base = {r["Service"]: r["OISI score strict all 4 (0-100)"] for r in complete}
    all_zero = {r["Service"]: r["Evidence-adjusted OISI (missing=0) (0-100)"] for r in rows}
    all_available = {r["Service"]: r["OISI score over available components (0-100)"] for r in rows if r["OISI score over available components (0-100)"] != ""}

    scenarios = []

    def add_scenario(name, cohort, method, ref, vals, notes):
        n, rho, top_overlap = rank_comparison(ref, vals)
        numeric = list(vals.values())
        scenarios.append({
            "Scenario": name,
            "Comparison cohort": cohort,
            "Method": method,
            "Entries scored": len(vals),
            "Compared entries": n,
            "Spearman rho vs reference": rho,
            "Top-10 overlap": top_overlap,
            "Mean score": round(sum(numeric) / len(numeric), 2) if numeric else "",
            "Minimum": round(min(numeric), 2) if numeric else "",
            "Maximum": round(max(numeric), 2) if numeric else "",
            "Interpretation": notes,
        })

    add_scenario("Reference: equal weights, complete cases", "53 complete cases", "25% each; no imputation needed", complete_base, complete_base, "Primary like-for-like reference ranking.")
    add_scenario("All 112: missing components = 0", "All 112", "Fixed 25% each; unavailable component contributes 0", all_zero, all_zero, "Secondary evidence-adjusted view; coverage must accompany the rank.")
    add_scenario("All 112: average available components", "All 112", "Equal average of available components", all_zero, all_available, "Diagnostic only; score can look high when few components are available.")

    # One-at-a-time perturbation: increase one component to 30%; share the 5-point reduction equally across the other three.
    for focus in range(4):
        weights = [25.0] * 4
        weights[focus] = 30.0
        for i in range(4):
            if i != focus:
                weights[i] = 70.0 / 3
        vals = {
            r["Service"]: round(sum(float(r[col]) * weight for col, weight in zip(score_columns, weights)) / 100, 2)
            for r in complete
        }
        label = f"{component_names[focus]} +5 percentage points"
        detail = f"Weights: " + "; ".join(f"{name} {weight:.2f}%" for name, weight in zip(component_names, weights))
        add_scenario(label, "53 complete cases", detail, complete_base, vals, "Stress test only; not a proposed replacement weight set.")

    gov_na = {}
    gov_na_score = {}
    for r in rows:
        g = gov_rows.get(norm(r["Service"]), {})
        alt_gov = governance_score_blank_sensitivity(g, blank_as_zero=False)
        gov_na_score[r["Service"]] = round(alt_gov, 2) if alt_gov is not None else ""
        parts = [r["Community Engagement score"], alt_gov, r["Organizational Capacity score"], r["Development Activity score"]]
        gov_na[r["Service"]] = round(sum(float(v or 0) for v in parts) / 4, 2)
    add_scenario("Governance blanks treated as missing", "All 112", "Recalculate governance with blank 048/050/046 excluded; outer missing score contributes 0", all_zero, gov_na, "Isolates the impact of blank-as-zero coding in the Governance source workbook.")

    detail_rows = []
    for r in rows:
        g = gov_rows.get(norm(r["Service"]), {})
        c = [float(r[col]) if r[col] != "" else None for col in score_columns]
        detail = {
            "Service": r["Service"],
            "Primary complete-case rank": r["Strict OISI rank"],
            "Primary strict score": r["OISI score strict all 4 (0-100)"],
            "All-112 zero-filled rank": r["Evidence-adjusted OISI rank"],
            "All-112 zero-filled score": r["Evidence-adjusted OISI (missing=0) (0-100)"],
            "Available-component average": r["OISI score over available components (0-100)"],
            "Components available": r["Components available (of 4)"],
            "Governance score, blank=0": r["Governance Practices score"],
            "Governance score, blank=missing": gov_na_score[r["Service"]],
            "Governance-blank sensitivity OISI": gov_na[r["Service"]],
        }
        if all(v is not None for v in c):
            for i, name in enumerate(component_names):
                weights = [25.0] * 4
                weights[i] = 30.0
                for j in range(4):
                    if j != i:
                        weights[j] = 70.0 / 3
                score = sum(float(r[col]) * w for col, w in zip(score_columns, weights)) / 100
                detail[f"{name} +5pp OISI"] = round(score, 2)
        else:
            for name in component_names:
                detail[f"{name} +5pp OISI"] = ""
        detail_rows.append(detail)
    rank_columns = {
        "Available-component rank": "Available-component average",
        "Governance-blank sensitivity rank": "Governance-blank sensitivity OISI",
        **{f"{name} +5pp rank": f"{name} +5pp OISI" for name in component_names},
    }
    for rank_column, score_column in rank_columns.items():
        values = {
            d["Service"]: float(d[score_column])
            for d in detail_rows
            if d.get(score_column) not in (None, "")
        }
        ranks = rank_map({name: -value for name, value in values.items()})
        for d in detail_rows:
            d[rank_column] = ranks.get(d["Service"], "")
    return scenarios, detail_rows


def build_equal_metric_results(weighted_rows):
    """Recalculate each component with equal weights across its currently scored metrics."""
    ce_records = by_name(read_table("Community-Engagement/OISI_Community_Engagement_v3_multi_sheet.xlsx", "Scoring"), "Service")
    gov_records = by_name(read_table("Governance-Practices/OISI_Governance_Practices_Scoring.xlsx", "Scores"), "Service")
    org_records = by_name(read_table("Organizational-capacity/OISI_Organizational_Capacity_Scoring_v3.xlsx", "Scores"), "Service")
    dev_records = by_name(read_table("Development-Activity/Dev Data/oisi_c4/OISI_Component4_Final_Master_With_Provenance.xlsx", "Scoring"), "infrastructure")

    metric_fields = {
        "Community Engagement": ["E1 Engagement availability", "E2 Engagement activity breadth", "E3 Engagement evidence strength", "C1 Contribution opportunity breadth", "C2 Contribution guidance", "C3 Open contribution infrastructure", "N1 Community norms / Code of Conduct"],
        "Governance Practices": ["s024", "s026", "s108", "s050", "s048", "s029", "s033", "s035", "s044", "s046"],
        "Organizational Capacity": ["M1 Technology readiness", "M2 Maintenance status", "M3 Infrastructure age", "I1 Formal organizational host", "I2 Affiliations / institutional support", "I3 SCOSS participation", "O1 Hosted/SaaS service availability", "O2 Registered service-provider ecosystem"],
        "Development Activity": ["m1_commits", "m2_active_weeks", "m3_recency", "m4_releases", "m5_contributors", "m6_effective_contributors", "m7_retention", "m8_issue_response", "m9_pr_response", "m10_issue_close90", "m11_pr_decision30", "m12_stars", "m13_forks"],
    }
    detail_rows = []
    for base in weighted_rows:
        name = base["Service"]
        ce = ce_records.get(norm(name), {})
        gov = gov_records.get(norm(name), {})
        org = org_records.get(norm(name), {})
        dev = dev_records.get(norm(name), {})
        values = {
            "Community Engagement": [num(ce.get(k)) for k in metric_fields["Community Engagement"]],
            "Governance Practices": [num(gov.get(k)) for k in metric_fields["Governance Practices"]],
            "Organizational Capacity": [num(org.get(k)) for k in metric_fields["Organizational Capacity"]],
            "Development Activity": [num(dev.get(k)) for k in metric_fields["Development Activity"]],
        }
        # Governance 100_board_level remains descriptive; the ten active score indicators are equally weighted.
        gov_groups = {
            "Formal": values["Governance Practices"][0:4],
            "Transparency": values["Governance Practices"][5:10],
        }
        gov_n = sum(v is not None for v in values["Governance Practices"])
        gov_score = None
        if gov_n / 10 >= 0.70 and any(v is not None for v in gov_groups["Formal"]) and any(v is not None for v in gov_groups["Transparency"]):
            gov_score = 100 * sum(v for v in values["Governance Practices"] if v is not None) / gov_n

        ce_groups = [values["Community Engagement"][0:3], values["Community Engagement"][3:6], values["Community Engagement"][6:7]]
        ce_n = sum(v is not None for v in values["Community Engagement"])
        ce_score = None
        if ce_n / 7 >= 0.60 and all(any(v is not None for v in group) for group in ce_groups):
            ce_score = 100 * sum(v for v in values["Community Engagement"] if v is not None) / ce_n

        saas = str(org.get("076 SaaS options") or "").strip().lower()
        provider_applicable = saas in {"through solution provider or third party vendor", "through third party vendor only"}
        applicable_org = [True] * 8
        if saas == "not applicable":
            applicable_org[6] = False
        if not provider_applicable:
            applicable_org[7] = False
        org_values = values["Organizational Capacity"]
        org_applicable_count = sum(applicable_org)
        org_n = sum(ok and val is not None for ok, val in zip(applicable_org, org_values))
        org_groups = [org_values[0:3], org_values[3:6], org_values[6:8]]
        org_group_present = sum(any(org_values[i] is not None and applicable_org[i] for i in indexes) for indexes in (range(0, 3), range(3, 6), range(6, 8)))
        org_score = None
        if org_applicable_count and org_n / org_applicable_count >= 0.70 and org_group_present >= 2:
            org_score = 100 * sum(val for ok, val in zip(applicable_org, org_values) if ok and val is not None) / org_n

        dev_values = values["Development Activity"]
        dev_n = sum(v is not None for v in dev_values)
        dim_indexes = [range(0, 4), range(4, 7), range(7, 11), range(11, 13)]
        dimension_ok = all(sum(dev_values[i] is not None for i in indexes) / len(indexes) >= 0.50 for indexes in dim_indexes)
        dev_score = None
        if dev_n / 13 >= 0.75 and dimension_ok:
            dev_score = sum(v for v in dev_values if v is not None) / dev_n

        component_scores = [ce_score, gov_score, org_score, dev_score]
        available = [v for v in component_scores if v is not None]
        equal_strict = sum(component_scores) / 4 if len(available) == 4 else None
        equal_available = sum(available) / len(available) if available else None
        equal_zero = sum(v or 0 for v in component_scores) / 4
        result = {
            "Service": name,
            "Provider": base["Provider"],
            "Solution category": base["Solution category"],
            "CE equal-metric score": round(ce_score, 2) if ce_score is not None else "",
            "CE metrics scored": ce_n,
            "CE metrics total": 7,
            "CE status": "Reportable" if ce_score is not None and ce_n / 7 >= 0.85 else ("Provisional" if ce_score is not None else "Not scored"),
            "Governance equal-metric score": round(gov_score, 2) if gov_score is not None else "",
            "Governance metrics scored": gov_n,
            "Governance metrics total": 10,
            "Governance status": "Scored" if gov_score is not None else "Not scored",
            "Org equal-metric score": round(org_score, 2) if org_score is not None else "",
            "Org metrics scored": org_n,
            "Org metrics applicable": org_applicable_count,
            "Org status": "Reportable" if org_score is not None and org_n / org_applicable_count >= 0.85 else ("Provisional" if org_score is not None else "Not scored"),
            "Development equal-metric score": round(dev_score, 2) if dev_score is not None else "",
            "Development metrics scored": dev_n,
            "Development metrics total": 13,
            "Development status": "Scored" if dev_score is not None else (dev.get("score_status") or dev.get("master_record_status") or "Insufficient data"),
            "Equal-metric components available (of 4)": len(available),
            "Equal-metric strict OISI (0-100)": round(equal_strict, 2) if equal_strict is not None else "",
            "Equal-metric OISI missing=0 (0-100)": round(equal_zero, 2),
            "Equal-metric average over available components": round(equal_available, 2) if equal_available is not None else "",
            "Equal-metric missing components": "; ".join(name for name, value in zip(["Community Engagement", "Governance Practices", "Organizational Capacity", "Development Activity"], component_scores) if value is None),
            "Weighted current OISI (strict)": base["OISI score strict all 4 (0-100)"],
            "Weighted current OISI rank": base["Strict OISI rank"],
            "Weighted current OISI (all 112, missing=0)": base["Evidence-adjusted OISI (missing=0) (0-100)"],
            "Weighted current components available": base["Components available (of 4)"],
        }
        for component, fields in metric_fields.items():
            prefix = {"Community Engagement": "CE", "Governance Practices": "Gov", "Organizational Capacity": "Org", "Development Activity": "Dev"}[component]
            for field, value in zip(fields, values[component]):
                scale = 100 if component != "Development Activity" else 1
                result[f"{prefix} metric {field}"] = round(value * scale, 2) if value is not None else ""
        detail_rows.append(result)

    complete_equal = sorted([r for r in detail_rows if r["Equal-metric strict OISI (0-100)"] != ""], key=lambda r: r["Equal-metric strict OISI (0-100)"], reverse=True)
    equal_ranks = {r["Service"]: i for i, r in enumerate(complete_equal, 1)}
    all_equal = sorted(detail_rows, key=lambda r: r["Equal-metric OISI missing=0 (0-100)"], reverse=True)
    zero_ranks = {r["Service"]: i for i, r in enumerate(all_equal, 1)}
    available_sorted = sorted([r for r in detail_rows if r["Equal-metric average over available components"] != ""], key=lambda r: r["Equal-metric average over available components"], reverse=True)
    available_ranks = {r["Service"]: i for i, r in enumerate(available_sorted, 1)}
    for r in detail_rows:
        r["Equal-metric strict rank"] = equal_ranks.get(r["Service"], "")
        r["Equal-metric all-112 rank"] = zero_ranks[r["Service"]]
        r["Equal-metric available-average rank"] = available_ranks.get(r["Service"], "")

    comparisons = []
    component_labels = [
        ("Community Engagement", "Community Engagement score", "CE equal-metric score"),
        ("Governance Practices", "Governance Practices score", "Governance equal-metric score"),
        ("Organizational Capacity", "Organizational Capacity score", "Org equal-metric score"),
        ("Development Activity", "Development Activity score", "Development equal-metric score"),
    ]
    for label, weighted_col, equal_col in component_labels:
        weighted_map = {r["Service"]: float(r[weighted_col]) for r in weighted_rows if r[weighted_col] != ""}
        equal_map = {r["Service"]: float(r[equal_col]) for r in detail_rows if r[equal_col] != ""}
        n, rho, overlap = rank_comparison(weighted_map, equal_map)
        shared = set(weighted_map) & set(equal_map)
        comparisons.append({"Comparison": label, "Basis": "Entries scoreable under both methods", "Entries compared": n, "Spearman rank correlation": rho, "Top-10 overlap": overlap,
                           "Weighted mean": round(sum(weighted_map[k] for k in shared) / len(shared), 2) if shared else "",
                           "Equal-metric mean": round(sum(equal_map[k] for k in shared) / len(shared), 2) if shared else ""})
    weighted_strict = {r["Service"]: float(r["OISI score strict all 4 (0-100)"]) for r in weighted_rows if r["OISI score strict all 4 (0-100)"] != ""}
    equal_strict = {r["Service"]: float(r["Equal-metric strict OISI (0-100)"]) for r in detail_rows if r["Equal-metric strict OISI (0-100)"] != ""}
    n, rho, overlap = rank_comparison(weighted_strict, equal_strict)
    shared = set(weighted_strict) & set(equal_strict)
    comparisons.append({"Comparison": "Combined OISI, strict complete cases", "Basis": "Entries with all four component scores under both methods", "Entries compared": n, "Spearman rank correlation": rho, "Top-10 overlap": overlap,
                       "Weighted mean": round(sum(weighted_strict[k] for k in shared) / len(shared), 2) if shared else "",
                       "Equal-metric mean": round(sum(equal_strict[k] for k in shared) / len(shared), 2) if shared else ""})
    weighted_all = {r["Service"]: float(r["Evidence-adjusted OISI (missing=0) (0-100)"]) for r in weighted_rows}
    equal_all = {r["Service"]: float(r["Equal-metric OISI missing=0 (0-100)"]) for r in detail_rows}
    n, rho, overlap = rank_comparison(weighted_all, equal_all)
    comparisons.append({"Comparison": "Combined OISI, all 112 zero-filled", "Basis": "All infrastructures; unavailable component scores contribute zero", "Entries compared": n, "Spearman rank correlation": rho, "Top-10 overlap": overlap,
                       "Weighted mean": round(sum(weighted_all.values()) / len(weighted_all), 2),
                       "Equal-metric mean": round(sum(equal_all.values()) / len(equal_all), 2)})
    return detail_rows, comparisons


def build_xlsx(rows):
    wb = Workbook()
    wb.remove(wb.active)
    equal_rows, equal_comparisons = build_equal_metric_results(rows)
    for name in ["Codebook", "Summary", "Primary Ranking (53)", "All 112 Evidence View", "Equal Metric Scores", "Equal Metric Ranking", "Equal Metric All 112", "Equal Metric Comparison", "OISI Index", "Component Scores", "Scoring Rules", "Sensitivity", "Sensitivity Detail", "Data"]:
        wb.create_sheet(name)

    codebook = [
        {"Sheet": "Summary", "Purpose": "Overall OISI counts and score distribution."},
        {"Sheet": "Primary Ranking (53)", "Purpose": "Like-for-like ranking using current scores for infrastructures with all four components scored."},
        {"Sheet": "All 112 Evidence View", "Purpose": "All entries ranked by the current zero-filled evidence-adjusted score, with coverage and missingness shown."},
        {"Sheet": "Equal Metric Scores", "Purpose": "Equal metric weights within each component, with every metric score and both methods' component results."},
        {"Sheet": "Equal Metric Ranking", "Purpose": "Like-for-like ranking for infrastructures eligible under the equal-metric rules in all four components."},
        {"Sheet": "Equal Metric All 112", "Purpose": "All entries ranked under equal metric weights, with unavailable component scores contributing zero and coverage shown."},
        {"Sheet": "Equal Metric Comparison", "Purpose": "Rank correlation and top-ten overlap between the existing weighted approach and equal metric weights."},
        {"Sheet": "OISI Index", "Purpose": "One row per infrastructure with evidence-adjusted, strict, and available-component OISI scores."},
        {"Sheet": "Component Scores", "Purpose": "Component-level scores and statuses for all 112 infrastructures."},
        {"Sheet": "Scoring Rules", "Purpose": "Formula and interpretation rules for the full OISI index."},
        {"Sheet": "Sensitivity", "Purpose": "Rank stability under missing-data and component-weight scenarios."},
        {"Sheet": "Sensitivity Detail", "Purpose": "Per-infrastructure alternative scores for sensitivity scenarios."},
        {"Sheet": "Data", "Purpose": "Full joined OISI index data with repository/content summary."},
    ]
    add_table(wb["Codebook"], ["Sheet", "Purpose"], codebook)

    status_counts = Counter(r["Index status"] for r in rows)
    strict_scores = [r["OISI score strict all 4 (0-100)"] for r in rows if r["OISI score strict all 4 (0-100)"] != ""]
    adjusted_scores = [r["Evidence-adjusted OISI (missing=0) (0-100)"] for r in rows]
    summary = [
        {"Item": "Total infrastructures", "Value": len(rows)},
        {"Item": "Complete OISI index rows", "Value": status_counts["Complete OISI index"]},
        {"Item": "Partial OISI index rows", "Value": status_counts["Partial OISI index"]},
        {"Item": "Insufficient component data rows", "Value": status_counts["Insufficient component data"]},
        {"Item": "Strict OISI mean", "Value": round(sum(strict_scores) / len(strict_scores), 2) if strict_scores else ""},
        {"Item": "Strict OISI min", "Value": min(strict_scores) if strict_scores else ""},
        {"Item": "Strict OISI max", "Value": max(strict_scores) if strict_scores else ""},
        {"Item": "Primary comparison cohort", "Value": "53 infrastructures with all four component scores; use Strict OISI rank for like-for-like comparison."},
        {"Item": "All-112 evidence-adjusted cohort", "Value": "Secondary view. Unavailable component scores contribute zero and must be interpreted with component coverage and missing score notes."},
        {"Item": "Equal-metric perspective", "Value": "Each scored/applicable metric has equal influence within its component. Existing component weights remain 25% each in the combined index."},
        {"Item": "Equal-metric four-component entries", "Value": sum(r["Equal-metric strict OISI (0-100)"] != "" for r in equal_rows)},
        {"Item": "Equal-metric OISI mean (complete under equal rules)", "Value": round(sum(float(r["Equal-metric strict OISI (0-100)"]) for r in equal_rows if r["Equal-metric strict OISI (0-100)"] != "") / max(1, sum(r["Equal-metric strict OISI (0-100)"] != "" for r in equal_rows)), 2)},
        {"Item": "Equal-metric OISI range (complete under equal rules)", "Value": f"{min(float(r['Equal-metric strict OISI (0-100)']) for r in equal_rows if r['Equal-metric strict OISI (0-100)'] != ''):.2f} to {max(float(r['Equal-metric strict OISI (0-100)']) for r in equal_rows if r['Equal-metric strict OISI (0-100)'] != ''):.2f}"},
        {"Item": "Equal-metric OISI mean (all 112, missing=0)", "Value": round(sum(float(r["Equal-metric OISI missing=0 (0-100)"]) for r in equal_rows) / len(equal_rows), 2)},
        {"Item": "Equal-metric component scores available", "Value": "Community Engagement 88; Governance Practices 108; Organizational Capacity 109; Development Activity 72"},
        {"Item": "Evidence-adjusted OISI mean (all 112)", "Value": round(sum(adjusted_scores) / len(adjusted_scores), 2)},
        {"Item": "Evidence-adjusted OISI min (all 112)", "Value": min(adjusted_scores)},
        {"Item": "Evidence-adjusted OISI max (all 112)", "Value": max(adjusted_scores)},
        {"Item": "Evidence-adjusted formula", "Value": "(CE + Governance + Organizational Capacity + Development Activity) / 4; unavailable component scores contribute 0. See coverage and missing score notes."},
        {"Item": "Formula", "Value": "OISI = 0.25*Community Engagement + 0.25*Governance Practices + 0.25*Organizational Capacity + 0.25*Development Activity."},
    ]
    add_table(wb["Summary"], ["Item", "Value"], summary)

    primary_rows = sorted(
        [r for r in rows if r["OISI score strict all 4 (0-100)"] != ""],
        key=lambda r: r["Strict OISI rank"],
    )
    primary_headers = [
        "Primary rank (complete 4/4)", "Service", "Provider", "Solution category",
        "Community Engagement score", "Governance Practices score",
        "Organizational Capacity score", "Development Activity score",
        "Strict OISI score (0-100)", "Coverage", "Repository URL",
    ]
    primary_table = [
        {
            "Primary rank (complete 4/4)": r["Strict OISI rank"],
            "Service": r["Service"], "Provider": r["Provider"], "Solution category": r["Solution category"],
            "Community Engagement score": r["Community Engagement score"],
            "Governance Practices score": r["Governance Practices score"],
            "Organizational Capacity score": r["Organizational Capacity score"],
            "Development Activity score": r["Development Activity score"],
            "Strict OISI score (0-100)": r["OISI score strict all 4 (0-100)"],
            "Coverage": "4 of 4", "Repository URL": r["Repository URL"],
        }
        for r in primary_rows
    ]
    add_table(wb["Primary Ranking (53)"], primary_headers, primary_table)

    all_rows = sorted(rows, key=lambda r: r["Evidence-adjusted OISI rank"])
    all_headers = [
        "Evidence-adjusted OISI rank", "Service", "Provider", "Solution category",
        "Community Engagement score", "Governance Practices score",
        "Organizational Capacity score", "Development Activity score",
        "Evidence-adjusted OISI (missing=0) (0-100)", "Components available (of 4)",
        "Index status", "Missing components", "Missing score notes", "Repository URL",
    ]
    add_table(wb["All 112 Evidence View"], all_headers, all_rows)

    index_headers = [
        "Strict OISI rank",
        "Evidence-adjusted OISI rank",
        "Index row",
        "Service",
        "Provider",
        "Solution category",
        "Repository URL",
        "OISI score strict all 4 (0-100)",
        "Evidence-adjusted OISI (missing=0) (0-100)",
        "OISI score over available components (0-100)",
        "Components available (of 4)",
        "Index status",
        "Missing components",
        "Missing score notes",
        "Content summary",
    ]
    add_table(wb["OISI Index"], index_headers, rows)

    component_headers = [
        "Service",
        "Evidence-adjusted OISI rank",
        "Evidence-adjusted OISI (missing=0) (0-100)",
        "Components available (of 4)",
        "Missing components",
        "Missing score notes",
        "Community Engagement score",
        "Community Engagement status",
        "Governance Practices score",
        "Governance completeness",
        "Organizational Capacity score",
        "Organizational Capacity status",
        "Development Activity score",
        "Development Activity status",
        "Development classification",
        "Development repo used",
        "Development host",
        "Development available weight %",
    ]
    add_table(wb["Component Scores"], component_headers, rows)

    rules = [
        {"Rule": "Component weights", "Details": "Each of the four OISI components contributes 25% to the strict overall index."},
        {"Rule": "Strict OISI score", "Details": "Calculated only when all four component scores are available. Formula: average of the four component scores."},
        {"Rule": "Evidence-adjusted OISI", "Details": "For the all-112 view, average the four equally weighted component scores and substitute 0 for any unavailable component score. This is a conservative evidence-adjusted result, not an observed component score."},
        {"Rule": "Coverage display", "Details": "Always show components available (of 4), missing component names, and source score statuses beside the evidence-adjusted score. The 53 complete entries have full component coverage."},
        {"Rule": "Available-component score", "Details": "Diagnostic score averaged over available components only. This is not a replacement for the strict OISI score."},
        {"Rule": "Interpretation", "Details": "Zero imputation reduces the index when a component is unscored. It does not establish that the infrastructure has zero real-world performance in that component; use the missing score notes to interpret the result."},
        {"Rule": "Primary comparison", "Details": "Use the 53 complete cases for the primary like-for-like ranking. The zero-filled all-112 index is a secondary evidence-coverage view."},
        {"Rule": "Equal metric perspective", "Details": "Within each component, calculate the arithmetic mean of its scored/applicable metric scores; each current scoring metric receives equal influence. Retain current metric score anchors and component eligibility safeguards, translating metric-weight completeness to metric-count completeness. The combined index still weights the four components equally at 25% each."},
        {"Rule": "Equal metric completeness", "Details": "Community Engagement: at least 60% of its 7 metrics and at least one metric in each sub-component. Governance: at least 70% of 10 active scored metrics and Formal plus Transparency represented; 100_board_level remains descriptive/excluded. Organizational Capacity: at least 70% of applicable metrics and at least two sub-components. Development Activity: at least 75% of 13 metrics and at least half of each dimension's metric count."},
        {"Rule": "Equal metric effective shares", "Details": "Community Engagement: each metric 1/7 (14.29%); engagement and contribution groups each have 3/7 (42.86%), norms 1/7 (14.29%). Governance: each of 10 active metrics 10%; Formal 40%, Participatory 10%, Transparency 50%. Organizational Capacity: each applicable metric 1/8 (12.5% when all apply); sub-components 37.5%, 37.5%, 25%. Development Activity: each metric 1/13 (7.69%); dimensions reflect their metric counts: Continuity 30.77%, Contributors 23.08%, Responsiveness 30.77%, Adoption 15.38%."},
        {"Rule": "Community Engagement", "Details": "Uses v3 Community Engagement workbook metrics E1, E2, E3, C1, C2, C3, N1."},
        {"Rule": "Governance Practices", "Details": "Uses Governance Practices scoring workbook, including formal governance, participatory governance, and transparency metrics."},
        {"Rule": "Organizational Capacity", "Details": "Uses v3 Organizational Capacity workbook metrics M1-M3, I1-I3, O1-O2."},
        {"Rule": "Development Activity", "Details": "Uses Component 4 metrics m1-m13 and requires at least 75% available metric weight."},
    ]
    add_table(wb["Scoring Rules"], ["Rule", "Details"], rules)

    add_table(wb["Equal Metric Scores"], list(equal_rows[0]), equal_rows)
    equal_complete = sorted([r for r in equal_rows if r["Equal-metric strict OISI (0-100)"] != ""], key=lambda r: r["Equal-metric strict rank"])
    equal_rank_headers = [
        "Equal-metric strict rank", "Service", "Provider", "Solution category",
        "CE equal-metric score", "Governance equal-metric score", "Org equal-metric score", "Development equal-metric score",
        "Equal-metric strict OISI (0-100)", "Weighted current OISI (strict)", "Weighted current OISI rank",
    ]
    add_table(wb["Equal Metric Ranking"], equal_rank_headers, equal_complete)
    equal_all_sorted = sorted(equal_rows, key=lambda r: r["Equal-metric all-112 rank"])
    equal_all_headers = [
        "Equal-metric all-112 rank", "Service", "Provider", "Solution category",
        "CE equal-metric score", "Governance equal-metric score", "Org equal-metric score", "Development equal-metric score",
        "Equal-metric OISI missing=0 (0-100)", "Equal-metric components available (of 4)", "Equal-metric missing components",
        "CE metrics scored", "CE metrics total", "Governance metrics scored", "Governance metrics total",
        "Org metrics scored", "Org metrics applicable", "Development metrics scored", "Development metrics total",
    ]
    add_table(wb["Equal Metric All 112"], equal_all_headers, equal_all_sorted)
    add_table(wb["Equal Metric Comparison"], list(equal_comparisons[0]), equal_comparisons)

    sensitivity_rows, detail_rows = build_sensitivity(rows)
    sensitivity_headers = ["Scenario", "Comparison cohort", "Method", "Entries scored", "Compared entries", "Spearman rho vs reference", "Top-10 overlap", "Mean score", "Minimum", "Maximum", "Interpretation"]
    add_table(wb["Sensitivity"], sensitivity_headers, sensitivity_rows)
    detail_headers = list(detail_rows[0])
    add_table(wb["Sensitivity Detail"], detail_headers, detail_rows)

    data_headers = list(rows[0])
    add_table(wb["Data"], data_headers, rows)
    wb.save(OUT_XLSX)


def doc_p(text: str = "", style: str | None = None) -> str:
    style_xml = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f'<w:p>{style_xml}<w:r><w:t xml:space="preserve">{escape(str(text))}</w:t></w:r></w:p>'


def doc_table(rows):
    out = ['<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="0" w:type="auto"/></w:tblPr>']
    for row in rows:
        out.append("<w:tr>")
        for cell in row:
            out.append(f'<w:tc><w:tcPr><w:tcW w:w="1800" w:type="dxa"/></w:tcPr>{doc_p(cell)}</w:tc>')
        out.append("</w:tr>")
    out.append("</w:tbl>")
    return "".join(out)


def build_docx(rows):
    status_counts = Counter(r["Index status"] for r in rows)
    strict_rows = [r for r in rows if r["OISI score strict all 4 (0-100)"] != ""]
    top10 = sorted(strict_rows, key=lambda r: r["OISI score strict all 4 (0-100)"], reverse=True)[:10]
    adjusted_top10 = sorted(rows, key=lambda r: r["Evidence-adjusted OISI (missing=0) (0-100)"], reverse=True)[:10]
    adjusted_scores = [r["Evidence-adjusted OISI (missing=0) (0-100)"] for r in rows]
    sensitivity_rows, _ = build_sensitivity(rows)
    equal_rows, equal_comparisons = build_equal_metric_results(rows)
    equal_complete = [r for r in equal_rows if r["Equal-metric strict OISI (0-100)"] != ""]
    body = [
        doc_p("OISI Full Index for 112 Infrastructures", "Title"),
        doc_p("Combined Community Engagement, Governance Practices, Organizational Capacity, and Development Activity"),
        doc_p("Formula: OISI = 0.25 CE + 0.25 Governance + 0.25 Organizational Capacity + 0.25 Development Activity."),
        doc_p("Evidence-adjusted formula: (CE + Governance + Organizational Capacity + Development Activity) / 4; unavailable component scores contribute 0. This conservative score is reported with the number of scored components and missingness notes. A zero contribution is not a measured zero-performance result."),
        doc_p("Primary comparison: use the 53 infrastructures with all four component scores for the like-for-like ranking. The all-112 zero-filled ranking is a secondary evidence-coverage view."),
        doc_p("Summary", "Heading1"),
        doc_table(
            [
                ["Item", "Value"],
                ["Total infrastructures", len(rows)],
                ["Complete strict OISI rows", status_counts["Complete OISI index"]],
                ["Partial index rows", status_counts["Partial OISI index"]],
                ["Insufficient component data rows", status_counts["Insufficient component data"]],
                ["Evidence-adjusted OISI mean (all 112)", round(sum(adjusted_scores) / len(adjusted_scores), 2)],
                ["Evidence-adjusted OISI range", f"{min(adjusted_scores):.2f} to {max(adjusted_scores):.2f}"],
            ]
        ),
        doc_p("Top Complete OISI Scores (Primary Comparison)", "Heading1"),
        doc_table(
            [["Rank", "Service", "OISI", "CE", "Gov", "Org", "Dev"]]
            + [
                [
                    r["Strict OISI rank"],
                    r["Service"],
                    r["OISI score strict all 4 (0-100)"],
                    r["Community Engagement score"],
                    r["Governance Practices score"],
                    r["Organizational Capacity score"],
                    r["Development Activity score"],
                ]
                for r in top10
            ]
        ),
        doc_p("Top Evidence-Adjusted OISI Scores (Secondary All-112 View)", "Heading1"),
        doc_table(
            [["Rank", "Service", "Adjusted OISI", "Components", "Missing components"]]
            + [
                [
                    r["Evidence-adjusted OISI rank"],
                    r["Service"],
                    r["Evidence-adjusted OISI (missing=0) (0-100)"],
                    r["Components available (of 4)"],
                    r["Missing components"] or "None",
                ]
                for r in adjusted_top10
            ]
        ),
        doc_p("Sensitivity Analysis", "Heading1"),
        doc_table(
            [["Scenario", "Compared entries", "Spearman rho", "Top-10 overlap"]]
            + [
                [
                    r["Scenario"],
                    r["Compared entries"],
                    r["Spearman rho vs reference"],
                    r["Top-10 overlap"],
                ]
                for r in sensitivity_rows[1:]
            ]
        ),
        doc_p("Equal Metric Perspective", "Heading1"),
        doc_p("Alternate view: equal arithmetic mean across the scored/applicable metrics within each component. The four components remain at 25% each. Metric scales, scoring anchors, applicability, and eligibility safeguards are retained; Governance field 100 remains descriptive and excluded."),
        doc_p(f"Equal-metric scores are eligible for {sum(r['CE equal-metric score'] != '' for r in equal_rows)} Community Engagement, {sum(r['Governance equal-metric score'] != '' for r in equal_rows)} Governance Practices, {sum(r['Org equal-metric score'] != '' for r in equal_rows)} Organizational Capacity, and {sum(r['Development equal-metric score'] != '' for r in equal_rows)} Development Activity entries. {len(equal_complete)} entries have all four equal-metric component scores; their mean OISI is {sum(float(r['Equal-metric strict OISI (0-100)']) for r in equal_complete) / max(1, len(equal_complete)):.2f}."),
        doc_table(
            [["Comparison", "Entries", "Spearman rho", "Top-10 overlap"]]
            + [[r["Comparison"], r["Entries compared"], r["Spearman rank correlation"], r["Top-10 overlap"]] for r in equal_comparisons]
        ),
        doc_p("All 112 Infrastructure Index Rows", "Heading1"),
        doc_table(
            [["Service", "Evidence-adjusted OISI", "Strict OISI", "Available Avg", "Components", "Missing components", "Status", "Repository / content"]]
            + [
                [
                    r["Service"],
                    r["Evidence-adjusted OISI (missing=0) (0-100)"],
                    r["OISI score strict all 4 (0-100)"],
                    r["OISI score over available components (0-100)"],
                    r["Components available (of 4)"],
                    r["Missing components"] or "None",
                    r["Index status"],
                    r["Development repo used"] or r["Repository URL"] or "not available",
                ]
                for r in rows
            ]
        ),
    ]
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:body>{"".join(body)}<w:sectPr><w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/><w:pgMar w:top="720" w:right="720" w:bottom="720" w:left="720"/></w:sectPr></w:body></w:document>'
    )
    with ZipFile(OUT_DOCX, "w", ZIP_DEFLATED) as z:
        z.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
</Types>""",
        )
        z.writestr(
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>""",
        )
        z.writestr(
            "word/_rels/document.xml.rels",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>""",
        )
        z.writestr(
            "word/styles.xml",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:rPr><w:sz w:val="18"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:rPr><w:b/><w:sz w:val="34"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:rPr><w:b/><w:sz w:val="28"/></w:rPr></w:style>
  <w:style w:type="table" w:styleId="TableGrid"><w:name w:val="Table Grid"/><w:tblPr><w:tblBorders><w:top w:val="single" w:sz="4"/><w:left w:val="single" w:sz="4"/><w:bottom w:val="single" w:sz="4"/><w:right w:val="single" w:sz="4"/><w:insideH w:val="single" w:sz="4"/><w:insideV w:val="single" w:sz="4"/></w:tblBorders></w:tblPr></w:style>
</w:styles>""",
        )
        z.writestr("word/document.xml", xml)


def main():
    rows = build_rows()
    build_xlsx(rows)
    build_docx(rows)
    print(OUT_XLSX)
    print(OUT_DOCX)
    print(Counter(r["Index status"] for r in rows))


if __name__ == "__main__":
    main()
