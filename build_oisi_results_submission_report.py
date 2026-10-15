from __future__ import annotations

import math
import statistics
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

import build_methodology_oisi_updated as doc


ROOT = Path(__file__).resolve().parent
INPUT = ROOT / "OISI_Full_Index_Equal_Component_Weight_112_Missing_Data_Views_Literature_Supported.xlsx"
OUTPUT = ROOT / "OISI_Results_and_Evidence_Submission_Report.docx"
FIGDIR = ROOT / "OISI_Report_Visuals"
FIGDIR.mkdir(exist_ok=True)

COLORS = {
    "ink": "#24312D", "muted": "#5F716B", "grid": "#D8E1DD", "paper": "#FFFFFF",
    "green": "#327260", "blue": "#3B7EA1", "gold": "#D6A43A", "coral": "#C65D52",
    "teal": "#6A9C91", "pale": "#EAF1EE", "slate": "#70847D", "lightblue": "#E7F0F5",
}
COMPONENTS = ["Community Engagement", "Governance Practices", "Organizational Capacity", "Development Activity"]
COMP_COLORS = [COLORS["green"], COLORS["blue"], COLORS["gold"], COLORS["coral"]]
METRIC_LABELS = {
    "E1": "Engagement status", "E2": "Engagement activity breadth", "C1": "Contribution opportunity breadth",
    "C2": "Contribution guidance", "C3": "Public contribution repository", "N1": "Code of Conduct status",
    "G1": "Governance structure status", "G2": "Bylaws or formal rules", "G3": "Board/steering-body presence",
    "G4": "Web accessibility status", "G5": "Equity commitment status", "G6": "Privacy commitment status",
    "G7": "Community governance formality", "M1": "Technology readiness", "M2": "Maintenance status",
    "H1": "Human-resource capacity", "I2": "Institutional affiliation/support", "O1": "Hosted/service-delivery capacity",
    "m1_commits": "Human commits", "m2_active_weeks": "Active weeks", "m3_recency": "Commit recency",
    "m4_releases": "Release activity", "m5_contributors": "Active contributors",
    "m6_effective_contributors": "Effective contributors", "m7_retention": "Contributor retention",
    "m8_issue_response": "Issue response <=14 days", "m9_pr_response": "PR response <=7 days",
    "m10_issue_close90": "Issue closure <=90 days", "m11_pr_decision30": "PR decision <=30 days",
    "m12_stars": "Stars gained", "m13_forks": "Forks created",
}


def font(size, bold=False):
    path = r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"
    return ImageFont.truetype(path, size)


def chart_base(title, subtitle, size=(1600, 900)):
    im = Image.new("RGB", size, COLORS["paper"])
    dr = ImageDraw.Draw(im)
    dr.text((70, 40), title, font=font(39, True), fill=COLORS["ink"])
    dr.text((72, 96), subtitle, font=font(22), fill=COLORS["muted"])
    return im, dr


def save_hist(scores):
    im, dr = chart_base("Distribution of all-entry OISI scores", "Zero-filled equal-component scores; n = 112. Bars are 10-point intervals.")
    x0, y0, x1, y1 = 130, 210, 1510, 720
    counts = [sum(1 for s in scores if i <= s < i + 10) for i in range(0, 100, 10)]
    counts[-1] = sum(1 for s in scores if 90 <= s <= 100)
    maxc = max(counts)
    for tick in range(0, maxc + 1, max(1, math.ceil(maxc / 5))):
        y = y1 - tick / maxc * (y1 - y0)
        dr.line((x0, y, x1, y), fill=COLORS["grid"], width=2)
        dr.text((70, y - 13), str(tick), font=font(18), fill=COLORS["muted"])
    bw = (x1 - x0) / 10
    for i, c in enumerate(counts):
        left = x0 + i * bw + 9
        top = y1 - c / maxc * (y1 - y0)
        dr.rounded_rectangle((left, top, left + bw - 18, y1), radius=8, fill=COLORS["green"])
        dr.text((left + bw / 2 - 10, top - 30), str(c), font=font(18, True), fill=COLORS["ink"])
        dr.text((left, y1 + 20), f"{i*10}-{i*10+9}", font=font(16), fill=COLORS["muted"])
    dr.text((x0, y1 + 65), "OISI score", font=font(19, True), fill=COLORS["ink"])
    im.save(FIGDIR / "01_oisi_distribution.png")


def save_component_profile(means, medians):
    im, dr = chart_base("Component score profile", "Zero-filled component scores, each component weighted equally in OISI.")
    x0, x1, y0 = 610, 1485, 235
    for tick in range(0, 101, 20):
        x = x0 + (x1 - x0) * tick / 100
        dr.line((x, y0 - 35, x, 735), fill=COLORS["grid"], width=2)
        dr.text((x - 13, 750), str(tick), font=font(18), fill=COLORS["muted"])
    for i, c in enumerate(COMPONENTS):
        y = y0 + i * 125
        dr.text((70, y + 14), c, font=font(24, True), fill=COLORS["ink"])
        width = (x1 - x0) * means[c] / 100
        dr.rounded_rectangle((x0, y, x0 + width, y + 38), radius=8, fill=COMP_COLORS[i])
        medx = x0 + (x1 - x0) * medians[c] / 100
        dr.line((medx, y - 5, medx, y + 47), fill=COLORS["ink"], width=4)
        dr.text((1290, y + 4), f"Mean {means[c]:.1f} | Med {medians[c]:.1f}", font=font(17), fill=COLORS["ink"])
    dr.text((x0, 795), "Score (0-100); bar = mean, vertical mark = median", font=font(19), fill=COLORS["muted"])
    im.save(FIGDIR / "02_component_profile.png")


def save_metric_coverage(metrics):
    im = Image.new("RGB", (1800, 1300), COLORS["paper"])
    dr = ImageDraw.Draw(im)
    dr.text((60, 35), "Observed data coverage by metric", font=font(38, True), fill=COLORS["ink"])
    dr.text((62, 90), "Observed flag from Metric Detail; coverage = observed entries / 112.", font=font(22), fill=COLORS["muted"])
    x0, x1, y0 = 400, 1590, 165
    for tick in (0, 25, 50, 75, 100):
        x = x0 + (x1 - x0) * tick / 100
        dr.line((x, y0 - 15, x, 1210), fill=COLORS["grid"], width=2)
        dr.text((x - 15, 1220), str(tick), font=font(17), fill=COLORS["muted"])
    group_last = None
    y = y0
    for comp, code, observed in metrics:
        if comp != group_last:
            if group_last is not None:
                y += 8
            dr.text((60, y + 3), comp, font=font(21, True), fill=COLORS["ink"])
            y += 30
            group_last = comp
        pct = 100 * observed / 112
        i = COMPONENTS.index(comp)
        dr.text((80, y + 3), code, font=font(18, True), fill=COLORS["muted"])
        dr.rounded_rectangle((x0, y, x0 + (x1 - x0) * pct / 100, y + 20), radius=5, fill=COMP_COLORS[i])
        dr.text((x1 + 13, y - 2), f"{observed}/112  ({pct:.0f}%)", font=font(16), fill=COLORS["ink"])
        y += 27
    dr.text((x0, 1260), "Observed coverage (%)", font=font(17), fill=COLORS["muted"])
    im.save(FIGDIR / "03_metric_coverage.png")


def save_coverage_scatter(rows):
    im, dr = chart_base("Observed metric count and zero-filled OISI", "Each point is one infrastructure; n = 112. The relationship is descriptive, not causal.")
    x0, x1, y0, y1 = 130, 1510, 720, 180
    for t in range(0, 101, 20):
        y = y1 + (y0 - y1) * (100 - t) / 100
        dr.line((x0, y, x1, y), fill=COLORS["grid"], width=2)
        dr.text((75, y - 12), str(t), font=font(18), fill=COLORS["muted"])
    for t in range(0, 32, 5):
        x = x0 + (x1 - x0) * t / 31
        dr.line((x, y1, x, y0), fill=COLORS["grid"], width=1)
        dr.text((x - 12, y0 + 20), str(t), font=font(17), fill=COLORS["muted"])
    for observed, score, _ in rows:
        x = x0 + (x1 - x0) * observed / 31
        y = y0 - (y0 - y1) * score / 100
        dr.ellipse((x - 5, y - 5, x + 5, y + 5), fill=COLORS["blue"], outline=COLORS["paper"])
    dr.text((x0, y0 + 65), "Observed metric count (out of 31)", font=font(19, True), fill=COLORS["ink"])
    dr.text((18, 145), "OISI", font=font(18, True), fill=COLORS["ink"])
    im.save(FIGDIR / "04_coverage_vs_oisi.png")


def save_rank_extremes(top, bottom):
    im, dr = chart_base("Highest and lowest all-entry OISI scores", "Zero-filled index; compare alongside coverage before interpreting differences.", size=(1600, 1020))
    x0, x1, y0 = 580, 1460, 205
    for tick in range(0, 101, 20):
        x = x0 + (x1 - x0) * tick / 100
        dr.line((x, y0 - 10, x, 810), fill=COLORS["grid"], width=1)
        dr.text((x - 12, 902), str(tick), font=font(16), fill=COLORS["muted"])
    dr.line((x0, 890, x1, 890), fill=COLORS["muted"], width=2)
    dr.text((55, 145), "Top 10", font=font(25, True), fill=COLORS["green"])
    dr.text((55, 510), "Bottom 10", font=font(25, True), fill=COLORS["coral"])
    for block, entries in enumerate([top, bottom]):
        base = y0 + block * 365
        for i, row in enumerate(entries):
            y = base + i * 31
            score = row["score"]
            label = row["service"][:42]
            dr.text((55, y + 2), f"{row['rank']:>3}. {label}", font=font(16), fill=COLORS["ink"])
            dr.rounded_rectangle((x0, y, x0 + (x1 - x0) * score / 100, y + 21), radius=4,
                                 fill=COLORS["green"] if block == 0 else COLORS["coral"])
            dr.text((x1 + 10, y + 1), f"{score:.1f}", font=font(16, True), fill=COLORS["ink"])
    im.save(FIGDIR / "05_rank_extremes.png")


def save_workflow():
    im = Image.new("RGB", (1600, 750), COLORS["paper"])
    dr = ImageDraw.Draw(im)
    dr.text((65, 40), "OISI construction and reporting workflow", font=font(38, True), fill=COLORS["ink"])
    boxes = [
        ("Source data", "Infra Finder v3.0\nZenodo\n+ R-based GitHub API collection", COLORS["green"]),
        ("Metric inputs", "Field mapping\nCategorical/free-text coding\n12-month activity measures", COLORS["blue"]),
        ("Component scores", "31 metrics\n0-100 scores\nEqual weights within component", COLORS["gold"]),
        ("OISI views", "4 components x 25%\nZero-filled all-entry\nComplete / available diagnostics", COLORS["coral"]),
    ]
    xs = [55, 455, 855, 1255]
    for i, (title, txt, color) in enumerate(boxes):
        x = xs[i]
        dr.rounded_rectangle((x, 235, x + 280, 500), radius=18, fill=COLORS["pale"], outline=color, width=5)
        dr.rounded_rectangle((x, 235, x + 280, 300), radius=18, fill=color)
        dr.text((x + 18, 252), title, font=font(23, True), fill=COLORS["paper"])
        for j, line in enumerate(txt.split("\n")):
            dr.text((x + 20, 340 + 44 * j), line, font=font(17), fill=COLORS["ink"])
        if i < 3:
            dr.line((x + 290, 368, x + 380, 368), fill=COLORS["muted"], width=6)
            dr.polygon([(x + 380, 368), (x + 357, 351), (x + 357, 385)], fill=COLORS["muted"])
    dr.text((60, 590), "Each stage produces an auditable transformation; source observations, metric scores and coverage flags remain available in the workbooks.", font=font(20), fill=COLORS["muted"])
    im.save(FIGDIR / "06_workflow.png")


def image_paragraph(path, rid, docpr, width_inches=6.3):
    wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    pic = "http://schemas.openxmlformats.org/drawingml/2006/picture"
    ET.register_namespace("wp", wp)
    ET.register_namespace("a", a)
    ET.register_namespace("pic", pic)
    im = Image.open(path)
    cx = int(width_inches * 914400)
    cy = int(cx * im.height / im.width)
    p = doc.el(doc.W, "p")
    pp = doc.sub(p, doc.W, "pPr")
    doc.sub(pp, doc.W, "jc", {f"{{{doc.W}}}val": "center"})
    r = doc.sub(p, doc.W, "r")
    drawing = doc.sub(r, doc.W, "drawing")
    inline = doc.sub(drawing, wp, "inline", {"distT": "0", "distB": "0", "distL": "0", "distR": "0"})
    doc.sub(inline, wp, "extent", {"cx": str(cx), "cy": str(cy)})
    doc.sub(inline, wp, "docPr", {"id": str(docpr), "name": path.name, "descr": path.stem.replace("_", " ")})
    frame = doc.sub(inline, wp, "cNvGraphicFramePr")
    doc.sub(frame, a, "graphicFrameLocks", {"noChangeAspect": "1"})
    graphic = doc.sub(inline, a, "graphic")
    gd = doc.sub(graphic, a, "graphicData", {"uri": pic})
    picture = doc.sub(gd, pic, "pic")
    nv = doc.sub(picture, pic, "nvPicPr")
    doc.sub(nv, pic, "cNvPr", {"id": "0", "name": path.name})
    doc.sub(nv, pic, "cNvPicPr")
    fill = doc.sub(picture, pic, "blipFill")
    doc.sub(fill, a, "blip", {f"{{{doc.R}}}embed": rid})
    stretch = doc.sub(fill, a, "stretch")
    doc.sub(stretch, a, "fillRect")
    sp = doc.sub(picture, pic, "spPr")
    xfrm = doc.sub(sp, a, "xfrm")
    doc.sub(xfrm, a, "off", {"x": "0", "y": "0"})
    doc.sub(xfrm, a, "ext", {"cx": str(cx), "cy": str(cy)})
    geom = doc.sub(sp, a, "prstGeom", {"prst": "rect"})
    doc.sub(geom, a, "avLst")
    return p


def add_figure(image, caption, n):
    rid = f"rId{n}"
    doc.body.append(image_paragraph(image, rid, n))
    doc.add(caption, style="Caption", align="center")


def build():
    wb = load_workbook(INPUT, data_only=True, read_only=True)
    summary = {wb["Summary"].cell(r, 1).value: wb["Summary"].cell(r, 2).value for r in range(2, wb["Summary"].max_row + 1)}
    ranking_headers = [c.value for c in wb["OISI Ranking"][1]]
    ranking = [dict(zip(ranking_headers, row)) for row in wb["OISI Ranking"].iter_rows(min_row=2, values_only=True)]
    complete_headers = [c.value for c in wb["Complete-Data Ranking"][1]]
    complete = [dict(zip(complete_headers, row)) for row in wb["Complete-Data Ranking"].iter_rows(min_row=2, values_only=True)]
    score_headers = [c.value for c in wb["Component Scores"][1]]
    component_rows = [dict(zip(score_headers, row)) for row in wb["Component Scores"].iter_rows(min_row=2, values_only=True)]
    metric_rows = list(wb["Metric Detail"].iter_rows(min_row=2, values_only=True))
    score_col = "Equal-component OISI (0-100)"
    scores = [float(r[score_col]) for r in ranking]
    means = {c: statistics.mean(float(r[f"{c} score (zero-filled)"]) for r in component_rows) for c in COMPONENTS}
    medians = {c: statistics.median(float(r[f"{c} score (zero-filled)"]) for r in component_rows) for c in COMPONENTS}
    observed_counts = {c: [r[10 + i * 4] if False else None for r in []] for i, c in enumerate(COMPONENTS)}
    metric_groups = {}
    for row in metric_rows:
        service, comp, code, score, observed = row[0], row[1], row[2], row[3], row[4]
        key = (comp, code)
        x = metric_groups.setdefault(key, {"scores_zero": [], "scores_obs": [], "n": 0})
        x["scores_zero"].append(float(score or 0))
        if observed == "Yes":
            x["n"] += 1
            x["scores_obs"].append(float(score or 0))
    metrics = [(c, code, metric_groups[(c, code)]["n"]) for c in COMPONENTS for comp, code in metric_groups if comp == c]
    oisi_by_service = {r["Service"]: r for r in ranking}
    rows_scatter = [(sum(int(component_rows[i][f"{c} observed metrics"]) for c in COMPONENTS), float(component_rows[i]["Equal-component OISI (0-100)"]), component_rows[i]["Service"]) for i in range(len(component_rows))]
    top = [{"rank": int(r["Rank (average ties)"]), "service": r["Service"], "score": float(r[score_col])} for r in ranking[:10]]
    bottom = [{"rank": int(r["Rank (average ties)"]), "service": r["Service"], "score": float(r[score_col])} for r in ranking[-10:]]

    save_hist(scores)
    save_component_profile(means, medians)
    save_metric_coverage(metrics)
    save_coverage_scatter(rows_scatter)
    save_rank_extremes(top, bottom)
    save_workflow()

    all_complete = int(summary["Complete-data comparison cases"])
    complete_mean = float(summary["Complete-data mean OISI"])
    available_mean = float(summary["Mean available-component diagnostic"])
    mean_oisi = float(summary["Mean OISI score (0-100)"])
    median_oisi = float(summary["Median OISI score (0-100)"])
    sd_oisi = statistics.stdev(scores)
    zeroed_means = [float(r["Equal-component OISI (0-100)"]) for r in ranking]
    q1, _, q3 = statistics.quantiles(scores, n=4, method="inclusive")
    no_missing_count = sum(r["All component metrics observed"] == "Yes" for r in ranking)
    all_metric_count = len(metric_groups)
    avg_coverage = {c: statistics.mean(float(r[f"{c} observed coverage"]) for r in component_rows) * 100 for c in COMPONENTS}
    mean_metric_counts = {c: statistics.mean(int(r[f"{c} observed metrics"]) for r in component_rows) for c in COMPONENTS}

    doc.body = []
    doc.add("OISI Results and Evidence", "Title", align="center", after=40)
    doc.add("Submission-ready results narrative, tables and figures", "Subtitle", align="center", after=180)
    doc.add("Equal-component index | 112 infrastructures | Four components", "Subtitle", align="center", after=220)
    doc.heading("Executive Results")
    doc.add(f"The all-entry equal-component OISI has a mean of {mean_oisi:.2f} (SD {sd_oisi:.2f}), median {median_oisi:.2f}, interquartile range {q1:.2f}–{q3:.2f}, and observed range {min(scores):.2f}–{max(scores):.2f} across 112 infrastructures. The four component means are Community Engagement {means[COMPONENTS[0]]:.2f}, Governance Practices {means[COMPONENTS[1]]:.2f}, Organizational Capacity {means[COMPONENTS[2]]:.2f}, and Development Activity {means[COMPONENTS[3]]:.2f}. Development Activity has the lowest mean in the current zero-filled view; this reflects measured scores and the availability of platform-specific data, so it should not be interpreted as a standalone verdict on infrastructure quality.")
    doc.add(f"Only {no_missing_count} infrastructures have all 31 metrics observed. Their complete-data mean is {complete_mean:.2f}; because this is a small subset, it is a diagnostic comparison, not a replacement for the all-entry result. The available-component diagnostic mean is {available_mean:.2f}, but its varying denominator makes it descriptive rather than a directly comparable performance ranking.")
    doc.add("The index is hierarchical: metrics are equally weighted within each component, and each component contributes 25% to OISI. Missing metrics are zero-filled in the all-entry score, with observation flags retained. Accordingly, report the result as an evidence-adjusted score and show component scores and coverage alongside it.")
    doc.tab(["Result item", "Value", "Reading"], [
        ["Infrastructures scored", "112", "All-entry OISI population."],
        ["All-entry OISI mean / median", f"{mean_oisi:.2f} / {median_oisi:.2f}", "Zero-filled; evidence-adjusted."],
        ["All-entry OISI SD / range", f"{sd_oisi:.2f} / {min(scores):.2f}–{max(scores):.2f}", "Describes cross-entry dispersion."],
        ["Complete across 31 metrics", f"{no_missing_count}", "Small complete-data comparison subset."],
        ["Complete-data mean", f"{complete_mean:.2f}", "Different cohort from all-entry mean."],
        ["Available-component diagnostic mean", f"{available_mean:.2f}", "Variable denominator; diagnostic only."],
    ], [3000, 1900, 4460])

    doc.heading("1. Analysis Basis")
    doc.add("This report summarizes the integrated equal-component index workbook and the four component scoring workbooks. The unit of analysis is one infrastructure/service. Community Engagement uses six metrics, Governance Practices seven, Organizational Capacity five, and Development Activity thirteen. All metric scores are on a 0–100 scale and are averaged within their component; the four component scores are then averaged with a fixed denominator of four.")
    doc.add("The Development Activity observation window is 1 October 2025 00:00 UTC inclusive to 1 October 2026 00:00 UTC exclusive. The first three components are based on the Infra Finder dataset fields; the fourth combines repository history and GitHub API-derived activity as defined in its scoring workbook.")
    doc.eq("Overall score", ["OISI = ", doc.mfrac("CE + G + OC + DA", "4"), " = 0.25CE + 0.25G + 0.25OC + 0.25DA"])
    doc.add("The workbook reports three views: the all-entry zero-filled OISI; a complete-data comparison restricted to entries with all 31 metrics observed; and an available-component diagnostic averaging observed metrics/components. These views answer different questions and should not be presented as interchangeable rankings.")
    add_figure(FIGDIR / "06_workflow.png", "Figure 1. Data-to-index workflow used to construct the reported results.", 1)

    doc.heading("2. Overall Score Distribution")
    doc.add(f"The zero-filled OISI spans {min(scores):.2f} to {max(scores):.2f}; its mean ({mean_oisi:.2f}) is below its median ({median_oisi:.2f}), indicating that lower-scoring entries pull the mean downward. Scores are distributed across the 0–100 scale rather than concentrated in a single narrow band. The histogram describes the observed index output, not an underlying normal population distribution.")
    add_figure(FIGDIR / "01_oisi_distribution.png", "Figure 2. Distribution of the all-entry zero-filled OISI (10-point bins). Source: integrated OISI workbook; authors' calculations.", 2)
    add_figure(FIGDIR / "05_rank_extremes.png", "Figure 3. Ten highest and ten lowest zero-filled OISI results. Interpret with metric coverage and component profiles.", 3)
    component_by_service = {r["Service"]: r for r in component_rows}
    doc.tab(["Rank", "Infrastructure", "OISI", "CE", "Governance", "Org. capacity", "Development", "Observed / 31"], [
        [int(r["Rank (average ties)"]), r["Service"], f"{float(r[score_col]):.2f}", f"{float(r['Community Engagement score']):.1f}", f"{float(r['Governance Practices score']):.1f}", f"{float(r['Organizational Capacity score']):.1f}", f"{float(r['Development Activity score']):.1f}", f"{sum(int(component_by_service[r['Service']][f'{c} observed metrics']) for c in COMPONENTS)}/31"]
        for r in ranking[:10]
    ], [620, 2280, 700, 700, 900, 1050, 900, 1210])
    doc.add("The top and bottom lists are descriptive rankings under the stated zero-fill assumptions. They do not establish causal explanations for score differences. Entries with sparse evidence may rank lower because missing metrics contribute zero.")

    doc.heading("3. Component Results")
    doc.add("Component means differ, with Governance Practices highest and Development Activity lowest. Since each component has a fixed 25% OISI share, the difference is not caused by the number of metrics in a component. The metric-level table in each subsection reports both the all-entry zero-filled mean and the observed-only mean, as well as observed count and coverage. The observed-only mean is conditional on the subset with source data and is not directly comparable to the zero-filled mean without noting that denominator.")
    add_figure(FIGDIR / "02_component_profile.png", "Figure 4. Component mean and median scores. Bars show means; vertical marks show medians.", 4)
    comp_headers = ["Component", "Mean", "Median", "Observed metrics / total (mean)", "Mean observed coverage"]
    comp_rows = [[c, f"{means[c]:.2f}", f"{medians[c]:.2f}", f"{mean_metric_counts[c]:.2f} / {dict(zip(COMPONENTS,[6,7,5,13]))[c]}", f"{avg_coverage[c]:.1f}%"] for c in COMPONENTS]
    doc.tab(comp_headers, comp_rows, [3100, 1200, 1200, 1900, 2060])

    doc.heading("3.1 Community Engagement", 2)
    doc.add(f"The Community Engagement zero-filled component mean is {means[COMPONENTS[0]]:.2f}, with a mean of {mean_metric_counts[COMPONENTS[0]]:.2f} of six metrics observed per entry ({avg_coverage[COMPONENTS[0]]:.1f}% average metric coverage). Contribution guidance and breadth measures have lower availability than the status and repository fields. The largest all-entry metric means are in public repository presence and engagement status, while breadth scores are lower, reflecting fewer supported activity/contribution types under the defined denominators.")
    add_metric_table("Community Engagement", metric_groups)
    doc.add("Submission interpretation: distinguish the availability of a public repository or policy/status response from the breadth and practical usability of community participation. URL matching is a scoring proxy, not an audit of link functionality.")

    doc.heading("3.2 Governance Practices", 2)
    doc.add(f"The Governance Practices zero-filled mean is {means[COMPONENTS[1]]:.2f}, with {mean_metric_counts[COMPONENTS[1]]:.2f} of seven metrics observed on average ({avg_coverage[COMPONENTS[1]]:.1f}% average coverage). Governance structure and board/body presence have relatively high all-entry metric means. Formal bylaws and reported commitments show more variation. The score represents reported structures and statuses rather than an assessment of their quality or effectiveness.")
    add_metric_table("Governance Practices", metric_groups)

    doc.heading("3.3 Organizational Capacity", 2)
    doc.add(f"The Organizational Capacity zero-filled mean is {means[COMPONENTS[2]]:.2f}, with {mean_metric_counts[COMPONENTS[2]]:.2f} of five metrics observed on average ({avg_coverage[COMPONENTS[2]]:.1f}% average coverage). Technology readiness and maintenance have high observed coverage and high all-entry means. Human-resource capacity and institutional affiliation/support have lower availability and substantially lower zero-filled means, so their contribution reflects both reported scores and gaps in records. Hosted/service-delivery scores measure reported availability, not service quality, uptime or affordability.")
    add_metric_table("Organizational Capacity", metric_groups)

    doc.heading("3.4 Development Activity", 2)
    doc.add(f"The Development Activity zero-filled mean is {means[COMPONENTS[3]]:.2f}, with {mean_metric_counts[COMPONENTS[3]]:.2f} of thirteen metrics observed on average ({avg_coverage[COMPONENTS[3]]:.1f}% average coverage). Git-derived continuity and recency measures are available for more entries than the issue-response and issue-closure rates, which require at least ten eligible events and a complete follow-up interval. Response and decision rates should therefore be read with their observed counts. Stars and forks are activity/attention proxies, not direct evidence of quality or successful downstream reuse.")
    add_metric_table("Development Activity", metric_groups)

    doc.heading("4. Metric Coverage and Missing Data")
    doc.add("Metric availability varies substantially. In this integrated workbook, the least observed fields include staffing/affiliation and event-rate measures that require sufficient GitHub activity; the most observed fields include technology readiness, maintenance, and several governance/status measures. A zero-filled value can be a measured zero or a policy zero assigned to unavailable data. The workbook's observed flags must accompany interpretations.")
    add_figure(FIGDIR / "03_metric_coverage.png", "Figure 5. Number of entries with an observed source value for each metric. The component colors follow the report palette.", 5)
    add_figure(FIGDIR / "04_coverage_vs_oisi.png", "Figure 6. Observed metric count versus zero-filled OISI for each infrastructure. This is descriptive and reflects the construction of the zero-filled score.", 6)
    doc.add("The plotted association between observed metric count and zero-filled score is partly built into the scoring rule: unavailable metrics are set to zero. It must not be interpreted as evidence that better reporting causes better infrastructure performance. For comparisons, report component values and coverage, and distinguish measured zero, zero-filled unavailable metric, and a missing component result.")

    doc.heading("5. Comparative Score Views")
    doc.tab(["Score view", "Cases", "Mean", "Use in reporting"], [
        ["All-entry zero-filled OISI", "112", f"{mean_oisi:.2f}", "Primary full-population evidence-adjusted view; reflects scores and evidence coverage."],
        ["Complete-data OISI", str(all_complete), f"{complete_mean:.2f}", "Strict all-31-metric comparison; small cohort, interpret cautiously."],
        ["Available-component diagnostic", "112 with >=1 component available", f"{available_mean:.2f}", "Descriptive variable-denominator diagnostic; not a comparable ranking."],
    ], [2450, 1600, 1300, 4110])
    doc.add("The complete-data mean should not be compared as if it represented all 112 entries: the six complete cases are a selected subset. The available-component mean is higher because it excludes unavailable metrics within components and averages only observed component results; it is not the same estimand as the fixed-denominator all-entry index. The three views are complementary reporting lenses, not substitutes.")

    doc.heading("6. Results Text for Submission")
    doc.add("The final integrated OISI included 112 infrastructures and combined four equally weighted components: Community Engagement, Governance Practices, Organizational Capacity, and Development Activity. Each component contributed 25% of the overall score, with equal metric weights within each component. In the all-entry zero-filled view, the mean OISI was " + f"{mean_oisi:.2f} (SD {sd_oisi:.2f}; median {median_oisi:.2f}; range {min(scores):.2f}–{max(scores):.2f}). Component means were " + "; ".join(f"{c}: {means[c]:.2f}" for c in COMPONENTS) + ". These are evidence-adjusted results because unavailable metric values were assigned zero while their observed status was retained.")
    doc.add(f"Six infrastructures had observed values for all 31 component metrics. Their mean score was {complete_mean:.2f}; this small complete-data subset is reported as a diagnostic and should not be generalized to the full population. The available-component diagnostic mean was {available_mean:.2f}; since it uses a variable denominator, it is descriptive rather than a directly comparable ranking. Metric-level availability differed across components and indicators, particularly for narrative staffing/support fields and GitHub event-rate measures. Results should therefore be interpreted alongside component scores and observed-data coverage.")

    doc.heading("7. Submission Checklist")
    doc.tab(["Report aspect", "Include in submission"], [
        ["Index construction", "Four component definitions, metric counts, hierarchical equal-weight rule, and overall formula."],
        ["Results", "All-entry mean, median, spread, component means, and clearly labeled top/bottom examples if ranking is reported."],
        ["Evidence coverage", "Metric observed counts/percentages and explicit explanation that zero-filled unavailable data are not confirmed absence."],
        ["Comparability", "Separate all-entry, complete-data, and available-component views; state different denominators and the six-case complete-data size."],
        ["Visual evidence", "Figures 1–6 with captions, source notes, readable labels and consistent score definitions."],
        ["Reproducibility", "Final workbook, data provenance, R extraction script and outputs, metric coding/rules, and observation-window dates."],
        ["Interpretation", "Describe OISI as a composite of selected indicators, not a validated measure of quality, impact or sustainability."],
    ], [2200, 7260])

    doc.heading("Appendix A. Full All-Entry Ranking")
    doc.add("All 112 entries are listed below using the zero-filled equal-component OISI. The observed count is the number of source-observed metrics out of 31. This appendix is intended to keep the report's ranking traceable; interpret rank alongside score composition and evidence coverage.")
    doc.tab(["Rank", "Infrastructure", "OISI", "CE", "Governance", "Org. capacity", "Development", "Observed /31"], [
        [int(r["Rank (average ties)"]), r["Service"], f"{float(r[score_col]):.2f}",
         f"{float(r['Community Engagement score']):.2f}", f"{float(r['Governance Practices score']):.2f}",
         f"{float(r['Organizational Capacity score']):.2f}", f"{float(r['Development Activity score']):.2f}",
         f"{sum(int(component_by_service[r['Service']][f'{c} observed metrics']) for c in COMPONENTS)}/31"]
        for r in ranking
    ], [620, 2280, 700, 700, 900, 1050, 900, 1210])
    doc.heading("Appendix B. Complete-Data Cases")
    doc.add("These six infrastructures have observed data for all 31 metrics. Their component observed-only scores and total are shown for strict complete-data comparison; do not generalize this small cohort's mean to all 112 entries.")
    doc.tab(["Complete rank", "Infrastructure", "OISI", "CE", "Governance", "Org. capacity", "Development"], [
        [r.get("Complete-data rank"), r.get("Service"), f"{float(r.get('Complete-data OISI (0-100)')):.2f}",
         f"{float(r.get('Community Engagement observed-only score')):.2f}",
         f"{float(r.get('Governance Practices observed-only score')):.2f}",
         f"{float(r.get('Organizational Capacity observed-only score')):.2f}",
         f"{float(r.get('Development Activity observed-only score')):.2f}"] for r in complete
    ], [1100, 2800, 1100, 1100, 1100, 1100, 1060])

    doc.heading("References")
    for i, text in enumerate([
        "Invest in Open Infrastructure. (2026). Data for Infra Finder, version 3.0. Zenodo record 18225266. https://zenodo.org/records/18225266",
        "Nardo, M., Saisana, M., Saltelli, A., Tarantola, S., Hoffmann, A., & Giovannini, E. (2008). Handbook on Constructing Composite Indicators: Methodology and User Guide. OECD/JRC. https://doi.org/10.1787/9789264043466-en",
        "Principles of Open Scholarly Infrastructure (POSI), version 1.0. https://doi.org/10.24343/C34W2H",
        "NASA. Technology Readiness Levels. https://www.nasa.gov/directorates/somd/space-communications-navigation-program/technology-readiness-levels/",
        "CHAOSS. Starter Project Health Metrics Model. https://www.chaoss.community/starter-project-health-metrics-model/",
        "Avelino, G., et al. (2019). On the abandonment and survival of open source projects: An empirical investigation. ESEM 2019. https://doi.org/10.1109/ESEM.2019.8870181",
        "Borges, H., & Valente, M. T. (2018). What's in a GitHub Star? Understanding Repository Starring Practices in a Social Coding Platform. Journal of Systems and Software. https://doi.org/10.1016/j.jss.2018.09.016",
        "Kalliamvakou, E., et al. (2014). The Promises and Perils of Mining GitHub. MSR 2014. https://doi.org/10.1145/2597073.2597074",
        "Yehudi, Y., Goble, C., & Jay, C. (2023). Individual context-free online community health indicators fail to identify open source software sustainability. arXiv:2309.12120. https://arxiv.org/abs/2309.12120",
    ], 1):
        doc.add(f"{i}. {text}")

    write_report_docx()
    print(f"Built report for {len(ranking)} entries and {all_metric_count} metrics; output={OUTPUT}")


def add_metric_table(component, groups):
    order = [code for comp, code in groups if comp == component]
    rows = []
    for code in order:
        info = groups[(component, code)]
        coverage = 100 * info["n"] / 112
        observed_mean = statistics.mean(info["scores_obs"]) if info["scores_obs"] else None
        rows.append([code, METRIC_LABELS.get(code, code), f"{info['n']}/112 ({coverage:.1f}%)",
                     f"{statistics.mean(info['scores_zero']):.2f}", f"{observed_mean:.2f}" if observed_mean is not None else "NA"])
    doc.tab(["Code", "Metric", "Observed / coverage", "All-entry mean (zero-filled)", "Observed-only mean"], rows,
            [1050, 2750, 1800, 1950, 1810])


def write_report_docx():
    wns = doc.W
    root = doc.el(wns, "document")
    body = doc.sub(root, wns, "body")
    for p in doc.body:
        body.append(p)
    sect = doc.sub(body, wns, "sectPr")
    doc.sub(sect, wns, "pgSz", {f"{{{wns}}}w": "12240", f"{{{wns}}}h": "15840"})
    doc.sub(sect, wns, "pgMar", {f"{{{wns}}}top": "850", f"{{{wns}}}right": "1440", f"{{{wns}}}bottom": "850", f"{{{wns}}}left": "1440", f"{{{wns}}}header": "420", f"{{{wns}}}footer": "420", f"{{{wns}}}gutter": "0"})
    content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/><Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/></Types>'''
    rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'''
    docrels = ['''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>''']
    image_paths = [FIGDIR / f for f in ["06_workflow.png", "01_oisi_distribution.png", "05_rank_extremes.png", "02_component_profile.png", "03_metric_coverage.png", "04_coverage_vs_oisi.png"]]
    for i, image_path in enumerate(image_paths, 2):
        docrels.append(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/{image_path.name}"/>')
    docrels.append("</Relationships>")
    core = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>OISI Results and Evidence</dc:title><dc:subject>Submission-ready index results and figures</dc:subject><dc:creator>OISI Research Team</dc:creator></cp:coreProperties>'''
    with ZipFile(OUTPUT, "w", ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/_rels/document.xml.rels", "".join(docrels))
        z.writestr("word/document.xml", ET.tostring(root, encoding="utf-8", xml_declaration=True))
        z.writestr("word/styles.xml", ET.tostring(doc.make_styles(), encoding="utf-8", xml_declaration=True))
        z.writestr("docProps/core.xml", core)
        for image_path in image_paths:
            z.write(image_path, f"word/media/{image_path.name}")


if __name__ == "__main__":
    build()
