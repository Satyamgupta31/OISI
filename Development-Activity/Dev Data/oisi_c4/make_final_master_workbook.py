"""Create final OISI Component 4 master workbook with provenance.

Inputs are the current pipeline CSV files in this directory. The workbook keeps
all 112 main-list entries and joins scoring, git, GitHub/API, resolved URL, and
recovery provenance data by oisi_id.
"""
import csv
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT = "OISI_Component4_Final_Master_With_Provenance.xlsx"


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def by_id(rows):
    return {r["oisi_id"]: r for r in rows}


def present(v):
    return str(v or "").strip() not in {"", "NA", "None", "nan"}


def safe_float(v):
    try:
        if not present(v):
            return None
        return float(v)
    except ValueError:
        return None


def source_note(path):
    p = Path(path)
    if not p.exists():
        return ""
    ts = datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
    return f"{path}; rows={len(read_csv(path))}; modified={ts}"


def add_table(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    style_sheet(ws)


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(10, min(42, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[letter].width = width


def status(row):
    if row.get("score_status") == "scored":
        return "Scored"
    if row.get("scored") != "yes":
        return "Not scored - excluded/not assessable"
    return "Insufficient data"


def main():
    repos = read_csv("repos.csv")
    git = by_id(read_csv("git_metrics.csv"))
    gh = by_id(read_csv("gh_metrics.csv"))
    scores = by_id(read_csv("scores.csv"))
    markers = by_id(read_csv("resolved_url_markers.csv"))
    recovery = by_id(read_csv("resolved_url_recovery_report.csv"))

    full = []
    for r in repos:
        oid = r["oisi_id"]
        joined = {}
        joined.update(r)
        for prefix, data in [
            ("git", git.get(oid, {})),
            ("gh", gh.get(oid, {})),
            ("score", scores.get(oid, {})),
            ("marker", markers.get(oid, {})),
            ("recovery", recovery.get(oid, {})),
        ]:
            for k, v in data.items():
                if k in {"oisi_id", "infrastructure"}:
                    continue
                key = k if prefix in {"git", "gh", "score"} and k not in joined else f"{prefix}_{k}"
                joined[key] = v
        joined["master_record_status"] = status(joined)
        full.append(joined)

    wb = Workbook()
    wb.remove(wb.active)

    tab_colors = {
        "Codebook": "006100",
        "Summary": "1F4E78",
        "Scoring Rules": "C00000",
        "Scoring": "548235",
        "Verification": "BF9000",
        "Calc": "7F7F7F",
        "Sensitivity": "7F7F7F",
        "Descriptive": "7F7F7F",
        "Parameters": "C00000",
        "Data": "7F7F7F",
    }

    for name in tab_colors:
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = tab_colors[name]

    # Codebook
    codebook_rows = [
        {"Sheet": "Summary", "Purpose": "Counts, coverage, and source-file provenance."},
        {"Sheet": "Scoring Rules", "Purpose": "Metric list, weights, formulas, eligibility rules, and reviewer-facing scoring interpretation."},
        {"Sheet": "Scoring", "Purpose": "Final Component 4 scores and metric scores for every main-list row."},
        {"Sheet": "Verification", "Purpose": "Repository decision, resolved URL markers, and recovery status for every main-list row."},
        {"Sheet": "Calc", "Purpose": "Raw metric inputs used by score.py, joined by oisi_id."},
        {"Sheet": "Sensitivity", "Purpose": "Weight sensitivity output from final_resolved_urls_sensitivity.csv."},
        {"Sheet": "Descriptive", "Purpose": "Descriptive metric availability counts and status distributions."},
        {"Sheet": "Parameters", "Purpose": "Metric definitions, weights, scoring anchors, and missing-data policy."},
        {"Sheet": "Data", "Purpose": "Full joined master table: repos + git + GitHub/API + score + provenance."},
    ]
    add_table(wb["Codebook"], ["Sheet", "Purpose"], codebook_rows)

    # Summary
    scored_counts = Counter(r["master_record_status"] for r in full)
    classifications = Counter(r.get("classification", "") for r in full)
    marker_counts = Counter(r.get("resolved_url_action", "") or "not in resolved list" for r in full)
    summary_rows = [
        {"Item": "Workbook generated UTC", "Value": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")},
        {"Item": "Main-list entries", "Value": len(repos)},
        {"Item": "Rows marked scoreable in repos.csv", "Value": sum(1 for r in repos if r.get("scored") == "yes")},
        {"Item": "Rows with complete git metrics", "Value": sum(1 for r in git.values() if r.get("git_status") == "complete")},
        {"Item": "Rows with GitHub/API metrics", "Value": len(gh)},
        {"Item": "Rows with final score", "Value": scored_counts["Scored"]},
        {"Item": "Rows insufficient data", "Value": scored_counts["Insufficient data"]},
        {"Item": "Rows excluded/not assessable", "Value": scored_counts["Not scored - excluded/not assessable"]},
        {"Item": "Resolved URL rows marked", "Value": sum(1 for r in full if r.get("resolved_url_marker") == "yes")},
        {"Item": "Resolved URL rows applied", "Value": marker_counts["applied"]},
        {"Item": "Resolved URL rows marker-only", "Value": marker_counts["marker_only"]},
        {"Item": "Input repos.csv", "Value": source_note("repos.csv")},
        {"Item": "Input git_metrics.csv", "Value": source_note("git_metrics.csv")},
        {"Item": "Input gh_metrics.csv", "Value": source_note("gh_metrics.csv")},
        {"Item": "Input scores.csv", "Value": source_note("scores.csv")},
        {"Item": "Input resolved_url_markers.csv", "Value": source_note("resolved_url_markers.csv")},
        {"Item": "Input resolved_url_recovery_report.csv", "Value": source_note("resolved_url_recovery_report.csv")},
    ]
    for key, count in sorted(classifications.items()):
        summary_rows.append({"Item": f"Classification: {key or '(blank)'}", "Value": count})
    add_table(wb["Summary"], ["Item", "Value"], summary_rows)

    # Scoring rules
    scoring_rule_rows = [
        {
            "Section": "Overall formula",
            "Rule": "Weighted average over available metric scores",
            "Details": "Component 4 score = score_on_available = weighted_points / available_weight_pct * 100, shown only when eligibility rules are met.",
        },
        {
            "Section": "Observation window",
            "Rule": "Fixed one-year window",
            "Details": "2025-10-01 00:00 UTC inclusive to 2026-10-01 00:00 UTC exclusive.",
        },
        {
            "Section": "Eligibility threshold",
            "Rule": "At least 75% of total metric weight must be available",
            "Details": "This is the main methodological threshold. In the current 112-entry master set, 72 repositories meet this threshold and receive scores.",
        },
        {
            "Section": "Dimension threshold",
            "Rule": "At least 50% of each dimension's weight must be available",
            "Details": "Prevents scoring a repository from only git metrics or only GitHub social/responsiveness metrics.",
        },
        {
            "Section": "Completeness finding",
            "Rule": "45 of the 72 scored repositories have 100% data availability",
            "Details": "These rows have all 13 metric scores available. Other scored rows meet the 75% threshold but have some unavailable metrics.",
        },
        {
            "Section": "Not scored finding",
            "Rule": "40 master entries do not receive a Component 4 score",
            "Details": "10 are excluded/not score-attempted (not-assessable, standards, or needs-decision); 30 are score-attempted but insufficient-data rows.",
        },
        {
            "Section": "Insufficient-data finding",
            "Rule": "Most insufficient rows fail because of low issue/PR volume or non-GitHub trackers",
            "Details": "Among the 30 insufficient rows, 24 have issue/PR counts below n >= 10 for responsiveness metrics and 6 are non-GitHub tracker/platform cases.",
        },
        {
            "Section": "m1 Human commits",
            "Rule": "Weight 10%; Continuity",
            "Details": "Count non-merge, non-bot commits in the window. Score = 100 * min(1, log10(1+c) / log10(1001)). Cap: 1000 commits.",
        },
        {
            "Section": "m2 Active weeks",
            "Rule": "Weight 15%; Continuity",
            "Details": "Distinct weeks with at least one human non-merge commit. Score = active_weeks / 52 * 100.",
        },
        {
            "Section": "m3 Recency",
            "Rule": "Weight 5%; Continuity",
            "Details": "100 if latest human commit is <=30 days before cutoff; linear decay to 0 at 365 days.",
        },
        {
            "Section": "m4 Release activity",
            "Rule": "Weight 5%; Continuity",
            "Details": "Git tags in window: 0=0, 1=50, 2-3=75, >=4=100.",
        },
        {
            "Section": "m5 Active contributors",
            "Rule": "Weight 10%; Contributors",
            "Details": "Resolved human commit identities with at least one non-merge commit. Score = 0 if n<=1, otherwise 100 * min(1, log2(n)/5). Cap: 32 contributors.",
        },
        {
            "Section": "m6 Effective contributors",
            "Rule": "Weight 10%; Contributors",
            "Details": "Inverse concentration of commit shares, E = 1 / sum(share^2). Score = 0 if E<=1, otherwise 100 * min(1, log2(E)/3). Cap: 8 effective contributors.",
        },
        {
            "Section": "m7 Contributor retention",
            "Rule": "Weight 5%; Contributors",
            "Details": "Contributors active in both halves of the window. Buckets: 0=0, 1=25, 2=50, 3-4=75, >=5=100.",
        },
        {
            "Section": "m8 Issue response",
            "Rule": "Weight 10%; Responsiveness",
            "Details": "Share of eligible issues receiving qualifying response within 14 days. At least 10 eligible issues required.",
        },
        {
            "Section": "m9 PR response",
            "Rule": "Weight 10%; Responsiveness",
            "Details": "Share of eligible PRs receiving qualifying response within 7 days. At least 10 eligible PRs required.",
        },
        {
            "Section": "m10 Issue closure within 90 days",
            "Rule": "Weight 7.5%; Responsiveness",
            "Details": "Share of eligible issues closed within 90 days. At least 10 eligible issues required.",
        },
        {
            "Section": "m11 PR decision within 30 days",
            "Rule": "Weight 7.5%; Responsiveness",
            "Details": "Share of eligible PRs merged or closed within 30 days. At least 10 eligible PRs required.",
        },
        {
            "Section": "m12 Stars gained",
            "Rule": "Weight 2.5%; Adoption",
            "Details": "GitHub stars gained in the window. Score = 100 * min(1, log10(1+s)/3). Kept equal with forks because both are indirect adoption proxies.",
        },
        {
            "Section": "m13 Forks created",
            "Rule": "Weight 2.5%; Adoption",
            "Details": "GitHub forks created in the window. Score = 100 * min(1, log10(1+f)/2). Kept equal with stars; sensitivity shows asymmetric reweighting has negligible effect.",
        },
    ]
    add_table(wb["Scoring Rules"], ["Section", "Rule", "Details"], scoring_rule_rows)

    # Scoring
    scoring_headers = [
        "oisi_id", "infrastructure", "classification", "scored", "master_record_status",
        "component4_score", "score_status", "available_weight_pct", "weighted_points", "score_on_available",
        "m1_commits", "m2_active_weeks", "m3_recency", "m4_releases", "m5_contributors",
        "m6_effective_contributors", "m7_retention", "m8_issue_response", "m9_pr_response",
        "m10_issue_close90", "m11_pr_decision30", "m12_stars", "m13_forks",
        "decision", "repo", "host", "issue_tracker", "notes",
    ]
    add_table(wb["Scoring"], scoring_headers, full)

    # Verification / provenance
    verification_headers = [
        "oisi_id", "infrastructure", "original_url", "classification", "decision", "scored",
        "host", "repo", "branch", "upstream", "issue_tracker", "issue_repo",
        "resolved_url_marker", "resolved_url_candidate_repo", "resolved_url_platform",
        "resolved_url_action", "resolved_url_note", "git_status", "gh_status",
        "recovery_git_status_before", "recovery_git_status_after",
        "recovery_available_weight_before", "recovery_available_weight_after",
        "recovery_remaining_gap", "notes",
    ]
    add_table(wb["Verification"], verification_headers, full)

    # Calc raw inputs
    calc_headers = [
        "oisi_id", "infrastructure", "repo", "git_status", "gh_status",
        "human_commits", "active_weeks", "days_since_last", "tags_in_window", "contributors",
        "eff_contributors", "retained", "m8_issue_response_rate", "m9_pr_response_rate",
        "m10_issue_close90_rate", "m11_pr_decision30_rate", "m12_stars_gained", "m13_forks_created",
        "active_months", "merge_commits", "bot_commits", "bot_share", "top_share", "bf50", "bf80",
        "diag_issues_created", "diag_prs_created", "diag_prs_merged", "diag_acceptance_ratio",
    ]
    add_table(wb["Calc"], calc_headers, full)

    # Sensitivity
    sens_path = Path("final_resolved_urls_sensitivity.csv")
    if sens_path.exists():
        sens_rows = read_csv(sens_path)
        add_table(wb["Sensitivity"], list(sens_rows[0]) if sens_rows else ["note"], sens_rows or [{"note": "No sensitivity rows"}])
    else:
        add_table(wb["Sensitivity"], ["note"], [{"note": "final_resolved_urls_sensitivity.csv not found"}])

    # Descriptive
    metric_cols = [
        "m1_commits", "m2_active_weeks", "m3_recency", "m4_releases", "m5_contributors",
        "m6_effective_contributors", "m7_retention", "m8_issue_response", "m9_pr_response",
        "m10_issue_close90", "m11_pr_decision30", "m12_stars", "m13_forks",
    ]
    desc_rows = []
    for m in metric_cols:
        vals = [safe_float(r.get(m)) for r in full if safe_float(r.get(m)) is not None]
        desc_rows.append({
            "metric": m,
            "available_rows": len(vals),
            "missing_rows": len(full) - len(vals),
            "min": round(min(vals), 2) if vals else "",
            "max": round(max(vals), 2) if vals else "",
            "average": round(sum(vals) / len(vals), 2) if vals else "",
        })
    for k, v in scored_counts.items():
        desc_rows.append({"metric": f"status: {k}", "available_rows": v, "missing_rows": "", "min": "", "max": "", "average": ""})
    add_table(wb["Descriptive"], ["metric", "available_rows", "missing_rows", "min", "max", "average"], desc_rows)

    # Parameters
    param_rows = [
        {"Parameter": "Observation window", "Value": "2025-10-01 00:00 UTC inclusive to 2026-10-01 00:00 UTC exclusive"},
        {"Parameter": "Common rules", "Value": "Default branch; author date; bots removed; merge commits excluded from authorship; identities resolved by rules.py."},
        {"Parameter": "Missing-data rule", "Value": "Score only when at least 75% total weight and at least half of each dimension weight are available; otherwise insufficient data."},
        {"Parameter": "m1 Human commits", "Value": "10%; non-merge non-bot commits; score=100*min(1,log10(1+c)/log10(1001))."},
        {"Parameter": "m2 Active weeks", "Value": "15%; distinct active weeks / 52 * 100."},
        {"Parameter": "m3 Recency", "Value": "5%; 100 if <=30 days, linear to 0 at 365 days."},
        {"Parameter": "m4 Release activity", "Value": "5%; 0=0, 1=50, 2-3=75, >=4=100."},
        {"Parameter": "m5 Active contributors", "Value": "10%; 100*min(1,log2(n)/5)."},
        {"Parameter": "m6 Effective contributors", "Value": "10%; 1/sum(share^2), score capped at 8 effective contributors."},
        {"Parameter": "m7 Contributor retention", "Value": "5%; contributors active in both half-years."},
        {"Parameter": "m8 Issue response rate", "Value": "10%; issue response within 14 days; >=10 eligible items required."},
        {"Parameter": "m9 PR response rate", "Value": "10%; PR response within 7 days; >=10 eligible items required."},
        {"Parameter": "m10 Issue closure within 90 days", "Value": "7.5%; >=10 eligible items required."},
        {"Parameter": "m11 PR decision within 30 days", "Value": "7.5%; >=10 eligible items required."},
        {"Parameter": "m12 Stars gained", "Value": "2.5%; GitHub star-history endpoint; score=100*min(1,log10(1+s)/3)."},
        {"Parameter": "m13 Forks created", "Value": "2.5%; forks created in window; score=100*min(1,log10(1+f)/2)."},
    ]
    add_table(wb["Parameters"], ["Parameter", "Value"], param_rows)

    # Full data
    preferred = [
        "oisi_id", "infrastructure", "master_record_status", "original_url", "classification", "decision", "scored",
        "host", "repo", "branch", "upstream", "issue_tracker", "issue_repo",
        "component4_score", "score_status", "available_weight_pct", "weighted_points", "score_on_available",
        "git_status", "gh_status", "resolved_url_marker", "resolved_url_candidate_repo",
        "resolved_url_platform", "resolved_url_action", "resolved_url_note",
    ]
    all_keys = []
    for row in full:
        for key in row:
            if key not in all_keys:
                all_keys.append(key)
    data_headers = preferred + [k for k in all_keys if k not in preferred]
    add_table(wb["Data"], data_headers, full)

    wb.save(OUT)
    print(OUT)
    print(f"rows={len(full)} scoreable={sum(1 for r in repos if r.get('scored') == 'yes')} scored={scored_counts['Scored']} insufficient={scored_counts['Insufficient data']}")


if __name__ == "__main__":
    main()
