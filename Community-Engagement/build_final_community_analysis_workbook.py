from __future__ import annotations

import importlib.util
import math
import statistics
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "OISI_Community_Engagement_v2_new_metrics.xlsx"
OUTPUT = HERE / "OISI_Community_Analysis_Final_Defensible.xlsx"

spec = importlib.util.spec_from_file_location("ce_metrics", HERE / "build_equal_weight_workbook.py")
ce = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(ce)

METRICS = [
    ("E1", "Engagement status", "037", "Engagement & Participation"),
    ("E2", "Engagement activity breadth", "038 + coded 938", "Engagement & Participation"),
    ("C1", "Contribution opportunity breadth", "040 + coded 940", "Contribution Mechanisms"),
    ("C2", "Contribution guidance", "041", "Contribution Mechanisms"),
    ("C3", "Public contribution repository", "019", "Contribution Mechanisms"),
    ("N1", "Code of Conduct status", "042; 043 used as evidence only", "Community Norms"),
]
WEIGHT = 1 / len(METRICS)


def style(ws, widths=None):
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="245B55")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        vals = [len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 80) + 1)]
        ws.column_dimensions[get_column_letter(col)].width = (widths or {}).get(
            col, min(max(max(vals, default=12) + 2, 12), 45)
        )


def metric_values(row):
    scores, notes, flags = ce.row_scores(row)
    return scores, notes, flags


def mean_available(values):
    vals = [v for v in values if isinstance(v, (int, float))]
    return sum(vals) / len(vals) if vals else None


def ordinal_ranks(score_map):
    ordered = sorted(score_map.items(), key=lambda kv: (-kv[1], kv[0]))
    return {name: i + 1 for i, (name, _) in enumerate(ordered)}


def main():
    source_wb = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)
    source_ws = source_wb["Data"]
    source_headers = [cell.value for cell in source_ws[1]]
    records = [dict(zip(source_headers, values)) for values in source_ws.iter_rows(min_row=2, values_only=True)]
    results = {}
    for row in records:
        results[row.get("001_service_name") or ""] = metric_values(row)

    observed_scores = {}
    zero_filled_scores = {}
    complete_scores = {}
    qualified_scores = {}
    for service, (vals, _, _) in results.items():
        observed = mean_available(vals)
        if observed is not None:
            observed_scores[service] = observed
        zero_filled_scores[service] = sum(v if isinstance(v, (int, float)) else 0 for v in vals) / 6
        if all(isinstance(v, (int, float)) for v in vals):
            complete_scores[service] = sum(vals) / 6
        n = sum(isinstance(v, (int, float)) for v in vals)
        dimension_coverage = (
            any(isinstance(v, (int, float)) for v in vals[:2])
            and any(isinstance(v, (int, float)) for v in vals[2:5])
            and isinstance(vals[5], (int, float))
        )
        if n >= 4 and dimension_coverage:
            qualified_scores[service] = observed

    complete_rank = ordinal_ranks(complete_scores)
    zero_rank = ordinal_ranks(zero_filled_scores)
    old_qualified_rank = ordinal_ranks(qualified_scores)
    zero_qualified_rank = ordinal_ranks({n: zero_filled_scores[n] for n in qualified_scores})
    rank_shifts = [abs(old_qualified_rank[n] - zero_qualified_rank[n]) for n in qualified_scores]
    rho = None
    if len(qualified_scores) > 1:
        xs = [old_qualified_rank[n] for n in qualified_scores]
        ys = [zero_qualified_rank[n] for n in qualified_scores]
        mx, my = statistics.mean(xs), statistics.mean(ys)
        denom = math.sqrt(sum((x - mx) ** 2 for x in xs) * sum((y - my) ** 2 for y in ys))
        rho = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / denom if denom else None
    top_observed = sorted(qualified_scores, key=lambda n: (-qualified_scores[n], n))[:10]
    top_zero_qualified = sorted(qualified_scores, key=lambda n: (-zero_filled_scores[n], n))[:10]
    top_overlap = len(set(top_observed) & set(top_zero_qualified))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    names = ["Summary", "Scoring", "Metric Rules", "Methodology", "Parameters", "Sensitivity", "Verification", "Activity Coding", "Contribution Coding", "Calc", "Review Queue", "Data"]
    sh = {name: wb.create_sheet(name) for name in names}

    sh["Metric Rules"].append(["Code", "Metric", "Source fields", "Weight", "Scoring criteria", "Unavailable-data treatment"])
    rule_rows = [
        ["E1", "Engagement status", "037", WEIGHT, "Implemented/available=100; In progress=67; Considering=33; Not planning=0.", "Unknown/blank remains unavailable in observed score; becomes 0 in zero-filled score."],
        ["E2", "Engagement activity breadth", "038 + coded 938", WEIGHT, "100 x distinct supported activity categories / 13. Twelve named categories plus valid Other; deduplicate across fields.", "No supported category is unavailable; becomes 0 in zero-filled score. 938 keyword coding is reviewable."],
        ["C1", "Contribution opportunity breadth", "040 + coded 940", WEIGHT, "100 x distinct supported contribution routes / 9. Eight named routes plus valid Other; deduplicate.", "No supported route is unavailable; becomes 0 in zero-filled score. Other requires selection in 040 and supporting 940 text."],
        ["C2", "Contribution guidance", "041", WEIGHT, "Dedicated guide URL pattern=100; other populated non-suspect URL=50.", "Blank or suspect URL unavailable; becomes 0 in zero-filled score."],
        ["C3", "Public contribution repository", "019", WEIGHT, "Recognized public forge host=100; other populated field=50.", "Blank unavailable; becomes 0 in zero-filled score. Host match is a proxy, not a live repository audit."],
        ["N1", "Code of Conduct status", "042; 043 evidence only", WEIGHT, "Implemented/available=100; In progress=67; Considering=33; Not planning=0.", "Unknown/blank unavailable; becomes 0 in zero-filled score. 043 is not separately scored."],
    ]
    for row in rule_rows:
        sh["Metric Rules"].append(row)

    method = sh["Methodology"]
    method.append(["Methodology and interpretation"])
    method_rows = [
        ["Metric weighting", "The six top-level metrics receive equal weights of 1/6 (16.67%) each. The equal weighting is a transparent design choice, not an empirically estimated claim that each construct has identical importance."],
        ["Primary comparison", f"Use the {len(complete_scores)} complete cases with all six metrics observed for the primary like-for-like performance ranking. This avoids treating missing fields as measured non-performance, while acknowledging that complete-case selection can introduce bias."],
        ["All-entry score", "The zero-filled score is provided for every infrastructure as a secondary evidence-adjusted result: each unavailable metric contributes 0 and the denominator remains six. It reflects evidence availability as well as measured metric values; do not describe it as performance alone."],
        ["Observed-only view", "The workbook also shows the mean of observed metric scores and the observed count (x/6). Because coverage varies, this view is descriptive and should not be used as a directly comparable full ranking."],
        ["Coverage-qualified sensitivity", f"A sensitivity subset of {len(qualified_scores)} entries has at least four observed metrics, at least one engagement metric, at least one contribution metric, and N1. Its observed-only and zero-filled rankings are compared on the same subset."],
        ["Missing-data distinction", "An explicit response of Not planning is scored 0 for E1/N1. Blank or Unknown is missing in the observed-only score. The secondary all-entry score intentionally imputes all unavailable metrics as 0; that assumption is visible and should not be confused with confirmed absence."],
        ["038/938 coding", "038 contains 12 named activity categories. A 13th Other category counts only when Other is selected and 938 text supports a distinct unlisted activity. 938 keyword coding is exposed in Activity Coding and Review Queue."],
        ["040/940 coding", "040 contains eight named contribution routes. A ninth Other route counts only when Other is selected and 940 text supports a distinct route. Not applicable is never counted as a route."],
        ["C2/C3 automation", "C2 uses URL text patterns to distinguish a likely dedicated contribution guide from another populated page. C3 uses host matching for public forges. These are reproducible proxies; the workbook does not claim live URL or repository verification."],
        ["Excluded fields", "116 is evidence only and not a separate metric. 048 Community Governance is not part of this Community Analysis score. 043 supports Code of Conduct review but earns no separate points."],
        ["No performance bands", "No 70-point threshold or high/medium/low score categories are applied."],
        ["Sensitivity evidence", f"On the same {len(qualified_scores)}-entry coverage-qualified subset, {sum(old_qualified_rank[n] != zero_qualified_rank[n] for n in qualified_scores)} ranks change between observed-only and zero-filled calculations; mean absolute shift={statistics.mean(rank_shifts):.2f}; Spearman rank correlation={rho:.3f}; top-10 overlap={top_overlap}/10." if rho is not None else "Sensitivity statistics unavailable."],
        ["Reference", "OECD, European Union, European Commission Joint Research Centre (2008). Handbook on Constructing Composite Indicators: Methodology and User Guide. See guidance on uncertainty/sensitivity analysis, including missing-data imputation and ranking robustness."],
        ["Reference URL", "https://doi.org/10.1787/9789264043466-en"],
    ]
    for item in method_rows:
        method.append(item)

    params = sh["Parameters"]
    params.append(["Code", "Metric", "Source fields", "Exact weight", "Displayed weight", "Dimension"])
    for code, name, fields, dimension in METRICS:
        params.append([code, name, fields, WEIGHT, "16.67%", dimension])
    params.append(["Total", "", "", sum(WEIGHT for _ in METRICS), "100.00%", ""])
    params.append([])
    params.append(["Breadth denominator", "Option count", "Count rule"])
    params.append(["038", 13, "12 named categories + supported Other"])
    params.append(["040", 9, "8 named routes + supported Other"])
    params.append(["Complete cases", len(complete_scores), "All six source-derived metric scores available"])
    params.append(["Coverage-qualified sensitivity cases", len(qualified_scores), "At least four metrics plus engagement, contribution, and N1 coverage"])

    scoring = sh["Scoring"]
    scoring.append(["Service", "Provider", "Solution category", "E1", "E2", "C1", "C2", "C3", "N1", "Observed metrics (of 6)", "Observed coverage", "Observed-only mean", "Complete-case?", "Primary complete-case score", "Primary complete-case rank", "All-entry zero-filled score", "All-entry zero-filled rank", "Coverage-qualified sensitivity?", "Data-quality notes"])
    verification = sh["Verification"]
    verification.append(["Service", "E1 override", "E2 override", "C1 override", "C2 override", "C3 override", "N1 override", "Reviewer notes"])
    calc = sh["Calc"]
    calc.append(["Service", "E1", "E2", "C1", "C2", "C3", "N1", "E2 category count", "C1 route count", "Data-quality flags"])
    review_q = sh["Review Queue"]
    review_q.append(["Service", "Review item", "Raw evidence / proposal", "Current treatment", "Reviewer notes"])

    act = sh["Activity Coding"]
    activity_categories = list(ce.ce.ENGAGEMENT_KEYWORDS)
    act.append(["Service", *[f"038: {c}" for c in activity_categories], *[f"938 coded: {c}" for c in activity_categories], "Union categories n", "Other counted", "E2 score", "Category union"])
    contribution = sh["Contribution Coding"]
    routes = ["Contribute funds", "Contribute to code", "Contribute to documentation", "Contribute to education or training", "Contribute to governance", "Contribute to peer reviews", "Contribute to user research or user testing", "Contribute to working groups or interest groups"]
    contribution.append(["Service", *[f"040: {route}" for route in routes], "940 text", "Other counted", "Route count", "C1 score", "Route union"])

    # Rows stay in source order; score/rank columns contain live Excel formulas.
    for idx, row in enumerate(records, start=2):
        service = row.get("001_service_name") or ""
        vals, notes, flags = results[service]
        calc.append([service, *vals, notes["activity_count"], notes["contribution_count"], flags])
        checklist, coded, union = ce.ce.engagement_categories(row.get("038_community_engagement_activities"), row.get("938_community_engagement_activities_other"))
        other_activity = "other" in ce.ce.lower(row.get("038_community_engagement_activities")) and bool(ce.ce.clean(row.get("938_community_engagement_activities_other")))
        activity_set = union | ({"Other"} if other_activity else set())
        act.append([service, *[int(c in checklist) for c in activity_categories], *[int(c in coded) for c in activity_categories], len(activity_set), int(other_activity), 100 * len(activity_set) / 13 if activity_set else None, "; ".join(sorted(activity_set))])

        selected_routes = {x for x in ce.ce.split_pipe(row.get("040_contribution_opportunities")) if x.lower() not in {"other", "not applicable", "unknown"}}
        other_route = "other" in {x.lower() for x in ce.ce.split_pipe(row.get("040_contribution_opportunities"))} and bool(ce.ce.clean(row.get("940_contribution_opportunities_other")))
        route_set = selected_routes | ({"Other"} if other_route else set())
        contribution.append([service, *[int(r in selected_routes) for r in routes], row.get("940_contribution_opportunities_other"), int(other_route), len(route_set), 100 * len(route_set) / 9 if route_set else None, "; ".join(sorted(route_set))])

        override_cols = "BCDEFG"
        metric_formulas = [
            f'=IF(ISNUMBER(Verification!{ocol}{idx}),Verification!{ocol}{idx},IF(Calc!{ccol}{idx}="","",Calc!{ccol}{idx}))'
            for ocol, ccol in zip(override_cols, override_cols)
        ]
        count_formula = f'=COUNT(D{idx}:I{idx})'
        coverage_formula = f'=J{idx}/6'
        observed_formula = f'=IFERROR(AVERAGE(D{idx}:I{idx}),"")'
        complete_formula = f'=IF(J{idx}=6,"Yes","No")'
        complete_score = f'=IF(M{idx}="Yes",L{idx},"")'
        complete_rank = f'=IF(N{idx}="","",RANK.EQ(N{idx},$N$2:$N${len(records)+1},0))'
        zero_score = f'=SUM(D{idx}:I{idx})/6'
        zero_rank = f'=RANK.EQ(P{idx},$P$2:$P${len(records)+1},0)'
        n = sum(isinstance(v, (int, float)) for v in vals)
        dimension_coverage = any(isinstance(v, (int, float)) for v in vals[:2]) and any(isinstance(v, (int, float)) for v in vals[2:5]) and isinstance(vals[5], (int, float))
        qualified = n >= 4 and dimension_coverage
        verification.append([service, None, None, None, None, None, None, "Use numeric override only to correct a source-derived metric score; blanks retain default."])
        scoring.append([service, row.get("008_provider_name"), row.get("014_solution_category"), *metric_formulas,
                        count_formula, coverage_formula, observed_formula, complete_formula, complete_score, complete_rank,
                        zero_score, zero_rank, "Yes" if qualified else "No", flags])

        if ce.ce.clean(row.get("938_community_engagement_activities_other")):
            review_q.append([service, "938 free-text activity coding", row.get("938_community_engagement_activities_other"), "; ".join(sorted(coded)), "Review category matches; Other only if distinct from all named types."])
        if ce.ce.clean(row.get("940_contribution_opportunities_other")):
            review_q.append([service, "940 Other contribution route", row.get("940_contribution_opportunities_other"), "Counted" if other_route else "Not counted", "Requires Other in 040 and a distinct route described in 940."])
        if ce.ce.clean(row.get("041_contribution_guidelines")):
            review_q.append([service, "041 contribution guidance", row.get("041_contribution_guidelines"), f'{notes["guidance_candidate_score"]}: {notes["guidance_note"]}', "URL pattern tier is automatic; inspect if clearly misclassified."])
        if ce.ce.clean(row.get("019_open_code_repository_url")):
            review_q.append([service, "019 repository field", row.get("019_open_code_repository_url"), f'{notes["repository_candidate_score"]}: {notes["repository_note"]}', "Host-based score is a proxy, not live verification."])
        if ce.ce.clean(row.get("043_code_of_conduct_url")):
            review_q.append([service, "043 Code of Conduct evidence", row.get("043_code_of_conduct_url"), "Evidence only", "043 is not a separate score."])

    summary = sh["Summary"]
    summary.append(["Community Analysis: final equal-metric results"])
    summary.append(["Infrastructure entries", len(records)])
    summary.append(["Metrics", "E1, E2, C1, C2, C3, N1"])
    summary.append(["Weight per metric", "1/6 = 16.67%"])
    summary.append(["Complete cases; primary comparable ranking", len(complete_scores)])
    summary.append(["Coverage-qualified sensitivity cases", len(qualified_scores)])
    summary.append(["All-entry zero-filled scores", len(records)])
    summary.append(["Mean complete-case score", '=IFERROR(AVERAGE(Scoring!N2:N%d),"")' % (len(records) + 1)])
    summary.append(["Mean all-entry zero-filled score", '=IFERROR(AVERAGE(Scoring!P2:P%d),"")' % (len(records) + 1)])
    summary.append(["Primary ranking", "Use the complete-case rank for like-for-like performance comparison; read Methodology for selection caveat."])
    summary.append(["Secondary ranking", "All-entry zero-filled score/rank includes a zero penalty for unavailable metrics and therefore also reflects evidence coverage."])
    summary.append(["Observed-only results", "Displayed with observed metric count and coverage; not a directly comparable rank where coverage differs."])
    summary.append(["Performance bands", "None; no 70-point cutoff."])

    sens = sh["Sensitivity"]
    sens.append(["Sensitivity summary", "Result"])
    sens.append(["Coverage-qualified entries", len(qualified_scores)])
    sens.append(["Entries whose within-cohort rank changes, observed-only vs zero-filled", sum(old_qualified_rank[n] != zero_qualified_rank[n] for n in qualified_scores)])
    sens.append(["Mean absolute rank shift", round(statistics.mean(rank_shifts), 2) if rank_shifts else "NA"])
    sens.append(["Spearman rank correlation", round(rho, 4) if rho is not None else "NA"])
    sens.append(["Top-10 overlap", f"{top_overlap}/10"])
    sens.append([])
    sens.append(["Service", "Observed-only score", "Zero-filled score", "Observed-only rank within coverage-qualified subset", "Zero-filled rank within same subset", "Rank shift", "Observed metric count"])
    for service in qualified_scores:
        sens.append([service, qualified_scores[service], zero_filled_scores[service], old_qualified_rank[service], zero_qualified_rank[service], zero_qualified_rank[service] - old_qualified_rank[service], sum(isinstance(v, (int, float)) for v in results[service][0])])

    raw = sh["Data"]
    raw.append(source_headers)
    for row in records:
        raw.append([row.get(h) for h in source_headers])

    for worksheet in sh.values():
        style(worksheet)
    custom_widths = {
        "Scoring": {1: 33, 2: 32, 3: 35, 20: 72},
        "Metric Rules": {1: 14, 2: 38, 3: 28, 5: 75, 6: 65},
        "Methodology": {1: 30, 2: 110},
        "Sensitivity": {1: 35, 2: 75, 3: 75, 4: 45, 5: 35, 6: 30, 7: 25},
        "Review Queue": {1: 32, 2: 34, 3: 80, 4: 45, 5: 80},
        "Verification": {1: 32, 8: 90},
        "Data": {i + 1: 36 for i in range(len(source_headers))},
    }
    for name, widths in custom_widths.items():
        for col, width in widths.items():
            sh[name].column_dimensions[get_column_letter(col)].width = width
    for cell in method[method.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("https://"):
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Entries={len(records)}, complete={len(complete_scores)}, qualified sensitivity={len(qualified_scores)}, zero-filled={len(records)}")
    print(f"Review queue rows={review_q.max_row-1}; same-cohort rank shifts={sum(old_qualified_rank[n] != zero_qualified_rank[n] for n in qualified_scores)}")


if __name__ == "__main__":
    main()
